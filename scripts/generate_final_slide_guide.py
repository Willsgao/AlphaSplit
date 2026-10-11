from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"F:\wills\competitions\AlphaSplit")
OUT_DIR = ROOT / "决赛答辩PPT_高玉伟三页_HTML"
OUTPUT = OUT_DIR / "AlphaSplit_决赛答辩三页内容与讲解说明.docx"
SLIDES = OUT_DIR / "rendered_slides"

NAVY = "17232D"
TEAL = "11B8AF"
BLUE = "3E8ED0"
AMBER = "E5A72B"
CORAL = "E86F61"
INK = "23313A"
MUTED = "5C6B73"
LINE = "CFD8DC"
PALE_TEAL = "EAF7F6"
PALE_BLUE = "EDF5FB"
PALE_AMBER = "FFF6E5"
PALE_CORAL = "FDEDEA"
PALE_GRAY = "F4F6F7"
WHITE = "FFFFFF"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs[edge]
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn("w:" + key), str(edge_data[key]))


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_headers(table):
    repeat_table_header(table.rows[0])


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def set_keep_lines(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepLines"))
    if keep is None:
        keep = OxmlElement("w:keepLines")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.append(begin)
    paragraph._p.append(instr)
    paragraph._p.append(end)
    paragraph.add_run(" 页")


def add_toc(paragraph):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在打开文档时自动更新"
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.extend([begin, instr, separate, end])


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(2.05)
    section.right_margin = Cm(2.05)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.22

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 12),
        ("Subtitle", 13, MUTED, 0, 12),
        ("Heading 1", 20, NAVY, 12, 8),
        ("Heading 2", 14, TEAL, 10, 5),
        ("Heading 3", 11.5, BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Cm(0.65)
        style.paragraph_format.first_line_indent = Cm(-0.35)
        style.paragraph_format.space_after = Pt(3)

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.text = "AlphaSplit  决赛答辩内容说明"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    set_keep_lines(p)
    return p


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(item, style=style)
        set_keep_lines(p)


def add_callout(doc, label, text, fill=PALE_TEAL, accent=TEAL):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.0)
    table.columns[1].width = Cm(13.7)
    left, right = table.rows[0].cells
    set_cell_fill(left, accent)
    set_cell_fill(right, fill)
    for cell in (left, right):
        set_cell_margin(cell, top=150, start=170, bottom=150, end=170)
        set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    r.font.size = Pt(10.5)
    r = right.paragraphs[0].add_run(text)
    r.font.size = Pt(10.3)
    prevent_row_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_fill(cell, NAVY)
        set_cell_margin(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.size = Pt(font_size)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    repeat_table_header(header)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        fill = WHITE if row_index % 2 == 0 else PALE_GRAY
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_fill(cell, fill)
            set_cell_margin(cell)
            set_cell_border(
                cell,
                top={"val": "single", "sz": 4, "color": LINE},
                bottom={"val": "single", "sz": 4, "color": LINE},
                left={"val": "single", "sz": 4, "color": LINE},
                right={"val": "single", "sz": 4, "color": LINE},
            )
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        numeric_widths = [float(width) for width in widths]
        total_width = sum(numeric_widths)
        scaled_widths = [width * 16.55 / total_width for width in numeric_widths]
        for i, width in enumerate(scaled_widths):
            table.columns[i].width = Cm(width)
        for row in table.rows:
            for i, width in enumerate(scaled_widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_equation(doc, text, caption=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, PALE_BLUE)
    set_cell_margin(cell, top=170, start=180, bottom=170, end=180)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 9, "color": BLUE},
        bottom={"val": "single", "sz": 9, "color": BLUE},
        left={"val": "single", "sz": 9, "color": BLUE},
        right={"val": "single", "sz": 9, "color": BLUE},
    )
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Cambria Math"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(12.5)
    if caption:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(caption)
        r2.font.size = Pt(8.8)
        r2.font.color.rgb = RGBColor.from_string(MUTED)
    prevent_row_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_slide_image(doc, filename, caption):
    path = SLIDES / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.52))
    cp = doc.add_paragraph(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(7)
    r = cp.runs[0]
    r.italic = True
    r.font.size = Pt(8.8)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_script_box(doc, time_label, text):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    top = table.cell(0, 0)
    body = table.cell(1, 0)
    set_cell_fill(top, NAVY)
    set_cell_fill(body, PALE_GRAY)
    set_cell_margin(top, top=90, start=170, bottom=90, end=170)
    set_cell_margin(body, top=170, start=190, bottom=170, end=190)
    p = top.paragraphs[0]
    r = p.add_run(f"现场讲解词  {time_label}")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    p = body.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(10.3)
    r.font.italic = True
    for row in table.rows:
        prevent_row_split(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_timing_table(doc, items):
    rows = []
    elapsed = 0
    for page, seconds, purpose in items:
        start = elapsed
        elapsed += seconds
        rows.append([page, purpose, f"{seconds}秒", f"{start}至{elapsed}秒"])
    add_table(doc, ["页码", "任务", "建议时长", "本段累计"], rows, [2.0, 8.6, 2.7, 3.4])


doc = Document()
configure_document(doc)

# Cover
section = doc.sections[0]
section.different_first_page_header_footer = True
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ALPHASPLIT")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor.from_string(TEAL)
r.font.name = "Aptos Display"

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("决赛答辩三页内容与讲解说明")

p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("第7页 可审计证据链  第8页 四层系统架构  第13页 6比4双轨治理")

doc.add_paragraph()
add_callout(
    doc,
    "核心主线",
    "证据如何产生  →  系统如何保证可信  →  证据如何进入治理",
    fill=PALE_TEAL,
    accent=TEAL,
)

for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("适用对象  高玉伟及决赛答辩制作团队\n").bold = True
p.add_run("内容版本  v1.0\n")
p.add_run("日期  2026年10月8日")
for run in p.runs:
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)

doc.add_page_break()
add_heading(doc, "目录", 1)
toc = doc.add_paragraph()
add_toc(toc)
add_callout(doc, "使用提示", "首次在 Word 中打开时，如目录页码未刷新，请在目录区域右键选择更新域。", fill=PALE_AMBER, accent=AMBER)

doc.add_page_break()
add_heading(doc, "文档用途与三页总逻辑", 1)
add_body(doc, "这三页不是三个孤立的技术点，而是一条从方法到系统、再到治理的闭环论证。第7页先证明贡献证据可以被规范地产生；第8页说明这条证据链为什么可信、可复算；第13页说明量化证据怎样进入机构分配决策，同时保留评审委员会的最终责任。")
add_table(
    doc,
    ["页码", "回答的问题", "核心结论", "评委应留下的印象"],
    [
        ["第7页", "角色贡献怎样算出来", "用统一口径完成32个联盟反事实回测，再以精确Shapley计量边际贡献", "不是主观打分，而是可复算证据"],
        ["第8页", "Agent会不会编造金融结果", "LLM只给受约束的结构化建议，金融数值由确定性引擎计算", "智能交互和金融事实有清晰边界"],
        ["第13页", "贡献证据怎样影响分配", "60%制度轨保连续性，40%证据轨做动态校正，评审委员会终裁", "证据进入治理，但不替代治理"],
    ],
    ["2.0", "4.2", "6.4", "4.1"],
)
add_heading(doc, "三页在八分钟答辩中的位置", 2)
add_body(doc, "决赛陈述总时长不超过8分钟。建议为这三页预留约100秒，占总陈述约21%。三页之间不要停顿解释旁支概念，应沿同一主线连续推进。")
add_timing_table(
    doc,
    [
        ("第7页", 40, "说明证据生成方法"),
        ("第8页", 35, "说明可信系统边界"),
        ("第13页", 25, "说明治理落地机制"),
    ],
)
add_heading(doc, "统一叙事句", 2)
add_callout(doc, "一句话", "AlphaSplit以更合理的分配机制为手段，把角色贡献转化为可审计证据，调动投研各角色参与积极性、促进协同与决策质量提升，并以长期投资回报改善为目标。", fill=PALE_BLUE, accent=BLUE)
add_heading(doc, "必须统一的角色边界", 2)
add_table(
    doc,
    ["主体", "可以做什么", "不能做什么"],
    [
        ["Agent", "理解自然语言、推荐白名单内因子、组织结构化回测计划、基于可信证据生成解释", "不能生成或改写收益率、风险指标、Shapley贡献和最终份额"],
        ["确定性金融引擎", "读取固定数据与配置，执行评分、调仓、成本计算、联盟回测和Shapley计量", "不能替代机构判断，也不自动执行资源划拨"],
        ["五类投研角色", "提供决策输入并基于证据陈述意见", "陈述不进入建议份额公式，也不是分配决策者"],
        ["评审委员会", "查阅证据、听取陈述、独立调整并确认方案", "不能把系统建议表述为自动裁决或收益承诺"],
    ],
    ["3.2", "7.0", "6.5"],
)

# Slide 07
doc.add_page_break()
add_heading(doc, "第7页 可审计证据链", 1)
add_body(doc, "页面标题  从角色决策到评审终裁，形成一条可审计证据链")
add_callout(doc, "页面任务", "把“角色贡献”从抽象管理问题，转化为可以输入、回测、计算、复核并进入治理的五步技术流程。", fill=PALE_TEAL, accent=TEAL)
add_slide_image(doc, "slide-07.png", "第7页内容原型  五步证据链与贯穿全程的审计记录")

add_heading(doc, "核心概念", 2)
add_table(
    doc,
    ["概念", "准确含义", "答辩时要避免的误解"],
    [
        ["角色建模", "把行业、个股、策略、量化、风控五类岗位视为合作博弈中的五个参与者", "不是把五个岗位简单变成五个自由对话的大模型"],
        ["因子代理", "用公开、可计算的因子表达角色的决策输入；当前为五角色乘以八因子，共40项输入", "因子是可运行代理，不等同于真实机构的全部决策日志"],
        ["角色联盟", "五角色任意组合形成的参与集合S；五个角色共有2的5次方即32个联盟", "不是只做删除单个角色的五次对照"],
        ["反事实回测", "只改变角色是否参与，其他数据、基准、成本和调仓口径保持一致，观察联盟结果变化", "不是为每个角色换一套样本或参数"],
        ["Shapley贡献", "角色在所有可能联盟中的加权平均边际贡献，用于公平分摊角色交互", "贡献可正可负；负值不代表程序出错"],
        ["可审计证据链", "数据版本、因子配置、联盟结果、贡献值、合成建议和人工确认均留痕", "可审计不等于保证未来收益"],
    ],
    ["3.0", "8.2", "5.5"],
)

add_heading(doc, "五步流程总览", 2)
add_table(
    doc,
    ["步骤", "输入", "处理", "输出", "审计点"],
    [
        ["一 角色建模", "投研职责与流程", "定义五类角色及职责边界", "角色集合N", "角色名称与版本"],
        ["二 因子代理", "客户偏好与因子库", "将偏好映射到白名单内的40项角色因子并人工确认", "生效因子配置", "因子权重与确认记录"],
        ["三 反事实回测", "固定数据、基准、成本、调仓规则", "枚举32个角色联盟并运行同口径Pipeline", "每个联盟的特征函数值v(S)", "数据版本与联盟结果"],
        ["四 贡献计量", "32个v(S)", "计算每个角色在全部联盟中的加权平均边际贡献", "带符号Shapley值", "可加性校验"],
        ["五 治理决策", "现行份额与证据池权重", "生成6比4合成建议并展示角色陈述", "供评审委员会调整确认的方案", "调整理由与人工确认"],
    ],
    ["2.7", "3.7", "5.0", "3.4", "2.7"],
    font_size=8.6,
)

add_heading(doc, "步骤一 角色建模", 3)
add_body(doc, "目的  明确“贡献归因给谁”。系统把行业、个股、策略、量化、风控五类角色定义为贡献计量对象，使组织职责可以进入合作博弈框架。")
add_bullets(doc, [
    "行业角色表达行业配置与行业判断输入。",
    "个股角色表达证券筛选与个股研究输入。",
    "策略角色表达择时、仓位或策略规则输入。",
    "量化角色表达多因子评分与组合构建输入。",
    "风控角色表达风险约束、暴露控制与风险投影输入。",
])
add_callout(doc, "输出标准", "输出必须是边界清晰的五角色集合N。后续任何联盟S都是N的子集，避免在回测过程中临时增删角色或改变职责。", fill=PALE_GRAY, accent=BLUE)

add_heading(doc, "步骤二 因子代理", 3)
add_body(doc, "目的  把难以直接计算的角色判断转化为可回测输入。当前系统为每类角色登记八项因子或参数，共40项输入，并把最终配置交给用户确认。")
add_bullets(doc, [
    "Agent读取客户偏好，只能从因子白名单中提出结构化建议。",
    "系统检查字段名称、参数范围、权重归一化和必填项。",
    "用户确认后，配置才成为回测输入；未确认的自然语言不直接驱动金融计算。",
    "机构落地时可用真实决策日志替换公开因子代理，但后续计量流程不变。",
])
add_callout(doc, "关键口径", "五角色乘以八因子描述的是40项可追溯输入，不是40个Agent，也不是声称所有机构都固定使用同一组因子。", fill=PALE_AMBER, accent=AMBER)

add_heading(doc, "步骤三 反事实回测", 3)
add_body(doc, "目的  构造公平的“有该角色”和“无该角色”对照。对任意角色联盟S，系统只启用S中角色的决策输入，在同一历史区间、同一HS300基准、同一交易成本、同一月度调仓规则下运行组合。")
add_equation(doc, "v(S) = 联盟S的组合收益 − 同期沪深300收益", "空联盟只持有基准，因此v(∅)=0；全集v(N)等于完整组合超额收益")
add_bullets(doc, [
    "五个角色共有32个不同联盟，系统进行全枚举。",
    "每次实验只改变角色参与集合，其他条件保持一致。",
    "联盟结果连同数据版本、成本假设和配置摘要一起留存。",
])

add_heading(doc, "步骤四 精确Shapley贡献计量", 3)
add_body(doc, "目的  把角色间的交互收益公平地分摊给各角色。仅看“全集减去单角色”的Leave One Out会依赖一个特定背景组合；Shapley值则对角色加入所有可能联盟时的边际贡献进行加权平均。")
add_equation(doc, "φᵢ = Σ  w(S) × [ v(S∪{i}) − v(S) ]", "求和遍历所有不含角色i的联盟S；w(S)按联盟规模给予标准Shapley权重")
add_bullets(doc, [
    "φᵢ大于零  角色在当前区间总体提升组合超额收益。",
    "φᵢ小于零  角色在当前区间总体形成拖累，仍应完整保留用于复盘。",
    "可加性校验  五个角色的Shapley值之和应等于全集超额收益v(N)。",
])

add_heading(doc, "步骤五 治理决策", 3)
add_body(doc, "目的  让量化证据进入治理，而不是让算法替代治理。带符号Shapley先转化为证据轨内部建议权重，再与机构现行份额按6比4合成。五角色可以陈述贡献依据，但其陈述不进入份额公式。评审委员会查阅证据后独立调整并确认。")
add_callout(doc, "权限边界", "系统输出的是可核验的合成建议，不自动分配预算、奖金、项目、数据或算力，也不承担评审委员会的最终决策责任。", fill=PALE_CORAL, accent=CORAL)

add_heading(doc, "审计链如何贯穿五步", 2)
add_table(
    doc,
    ["审计对象", "至少保留的内容", "复核目的"],
    [
        ["数据", "数据来源、时间区间、缓存版本、行业映射", "证明不同联盟使用同一输入口径"],
        ["配置", "角色、因子、权重、调仓频率、成本、基准、人工确认时间", "证明回测参数可重建"],
        ["执行", "32个联盟状态、运行日志、异常与降级记录", "证明没有选择性遗漏联盟"],
        ["结果", "联盟收益、完整组合指标、带符号Shapley与可加性校验", "证明贡献结果可复算"],
        ["治理", "现行份额、证据池权重、合成建议、角色陈述、人工调整与确认", "证明系统建议和最终决策有边界"],
    ],
    ["3.0", "8.7", "5.0"],
)

add_heading(doc, "页面视觉与动画顺序", 2)
add_bullets(doc, [
    "先出现标题和副标题，让评委先接受“统一口径”这一前提。",
    "五张流程卡从左到右依次出现，讲到哪一步就高亮哪一步。",
    "最后出现底部审计记录条和右侧Agent边界，收束为“可复算、非生成事实”。",
    "不要在此页展开6比4公式细节，只说“进入证据轨并由评审委员会确认”，把细节留到第13页。",
])

add_heading(doc, "现场讲解词", 2)
add_script_box(doc, "建议40秒", "我们的核心不是再做一套收益归因，而是把每个角色的决策输入转成一条可审计证据链。首先将行业、个股、策略、量化、风控五类角色映射到40项因子输入；人工确认配置后，在同一数据、同一成本、同一基准下，对32个角色联盟进行反事实回测；随后用精确Shapley得到每个角色带符号的边际贡献。最后，贡献证据只进入证据轨，与现行制度形成合成建议，再由评审委员会独立确认。整个链路的输入、版本和结果都能复算。")

add_heading(doc, "页间衔接", 2)
add_table(doc, ["方向", "衔接话术"], [
    ["承接上一页", "前面说明了为什么要把投研角色贡献从主观评价变成量化证据，下面看这条证据是怎样生成的。"],
    ["进入第8页", "这条链路要可信，关键是不能让大模型直接生成金融结果，因此我们在系统架构上做了明确分层。"],
], ["3.2", "13.5"])

add_heading(doc, "可能问答", 2)
add_table(doc, ["评委可能提问", "建议回答"], [
    ["为什么是32个联盟", "五个角色的任意子集共有2的5次方即32种。当前角色数较少，完整枚举可承受，因此使用精确Shapley而不是抽样近似。"],
    ["因子能代表真实角色吗", "当前公开因子是角色决策输入的可运行代理，用于验证方法和产品闭环；机构落地时可替换为脱敏后的真实决策日志，计量框架不变。"],
    ["负贡献是不是说明这个岗位没价值", "不是。它只说明在特定区间、特定配置下的平均边际作用为负，用于复盘而非永久评价，更不能直接等同于人员价值。"],
    ["历史回测能否证明未来收益", "不能。回测提供可复算的历史证据，不构成收益承诺；项目目标是改善协同和决策质量，长期促进投资回报。"],
], ["5.0", "11.7"])

# Slide 08
doc.add_page_break()
add_heading(doc, "第8页 四层系统架构", 1)
add_body(doc, "页面标题  AlphaSplit系统架构  智能交互、确定性计算与人工治理")
add_callout(doc, "页面任务", "回答可信性问题  Agent负责理解和组织，确定性金融引擎负责计算事实，人工治理负责最终确认。", fill=PALE_BLUE, accent=BLUE)
add_slide_image(doc, "slide-08.png", "第8页内容原型  四层架构与贯穿四层的审计脊柱")

add_heading(doc, "核心概念", 2)
add_table(doc, ["概念", "含义", "在AlphaSplit中的落点"], [
    ["分层架构", "把不同责任、数据和变更频率拆到独立层级", "数据与审计层、金融计算层、Agent编排层、交互与治理层"],
    ["确定性引擎", "同一输入和版本应得到同一可复算结果", "40因子评分、月度调仓、交易成本、32联盟回测、精确Shapley"],
    ["受约束Agent", "只能在白名单、参数范围和证据引用规则内生成结构化建议", "DeepSeek在线模式或离线缓存，失败时规则降级，提交前人工确认"],
    ["离线缓存", "预置并经过同样校验的Agent响应", "无API Key或无网络时仍能完整演示，不代表绕过校验"],
    ["审计脊柱", "输入、版本、执行和结果的追踪能力贯穿所有层", "从数据版本到人工确认均可查"],
], ["3.3", "6.4", "10.2"])

add_heading(doc, "四层输入处理输出", 2)
add_table(doc, ["层级", "主要输入", "核心处理", "主要输出", "可信边界"], [
    ["L1 数据与审计层", "HS300行情、行业映射、公开因子", "版本固定、缓存管理、月度诊断、日志记录", "可重建的数据快照与审计信息", "不生成投资结论"],
    ["L2 金融计算层", "固定数据与已确认配置", "评分、调仓、成本、32联盟回测、精确Shapley", "收益风险指标、联盟值、角色贡献", "不接受LLM直接写入结果"],
    ["L3 Agent编排层", "自然语言偏好、白名单、可信结果", "配置建议、参数校验、规划、证据约束解读、规则降级", "结构化计划与解释文本", "不生成金融事实，不终裁"],
    ["L4 交互与治理层", "客户偏好、结构化建议、回测证据、现行份额", "确认配置、展示证据、收集角色陈述、评审调整确认", "可应用的人工确认方案", "系统不自动执行划拨"],
], ["3.1", "3.7", "5.4", "3.8", "3.7"], font_size=8.3)

add_heading(doc, "层间调用顺序", 2)
add_bullets(doc, [
    "一  用户在交互层输入投资偏好，并选择在线Agent或离线完整演示。",
    "二  Agent层将自然语言转为因子白名单内的结构化配置建议。",
    "三  参数校验检查字段、类型、范围、归一化和必填项。",
    "四  用户在交互层确认配置；只有已确认配置可以进入回测。",
    "五  金融计算层从数据层读取固定版本的行情、行业和因子。",
    "六  引擎执行完整组合与32个角色联盟的同口径回测。",
    "七  引擎计算绩效指标、带符号Shapley贡献和可加性校验。",
    "八  结果写入审计记录，并通过证据白名单提供给Agent解释。",
    "九  交互层展示原始数值、图表、解释和角色陈述参考。",
    "十  评审委员会基于证据独立调整确认，系统记录最终决定。",
], numbered=False)

add_heading(doc, "在线与离线路径", 2)
add_table(doc, ["场景", "执行路径", "用户看到什么", "金融结果是否变化"], [
    ["配置了API Key", "自然语言 → DeepSeek → 结构化建议 → 本地校验 → 人工确认", "在线Agent建议与来源标识", "不会；金融结果仍由本地确定性引擎计算"],
    ["未配置API Key且命中缓存", "自然语言 → 预置Agent响应 → 本地校验 → 人工确认", "离线完整演示与缓存来源标识", "不会；缓存只提供配置建议"],
    ["未命中缓存或在线失败", "自然语言 → 规则Agent降级 → 本地校验 → 人工确认", "可解释的降级建议与提示", "不会；回测入口和计算逻辑不变"],
], ["3.6", "7.8", "4.2", "3.5"])
add_callout(doc, "关键澄清", "离线缓存缓存的是Agent处理后的结构化建议，不是伪造的金融结果。回测收益、风险和贡献仍由确定性引擎依据本地数据重新计算。", fill=PALE_AMBER, accent=AMBER)

add_heading(doc, "确定性计算与大模型的边界", 2)
add_table(doc, ["任务", "负责模块", "理由"], [
    ["理解客户偏好", "Agent", "自然语言存在表达差异，适合语义理解"],
    ["推荐因子配置", "Agent加白名单和校验", "需要交互灵活性，但必须受约束"],
    ["收益率与风险指标", "确定性金融引擎", "数值必须可重复计算"],
    ["32联盟反事实", "确定性金融引擎", "每个联盟必须使用同一口径"],
    ["Shapley贡献", "确定性金融引擎", "需要满足可加性等数学核验"],
    ["结果解释", "Agent读取证据白名单", "可提高可读性，但不得新增未经计算的数字"],
    ["最终分配确认", "评审委员会", "属于机构治理责任，不应外包给模型"],
], ["5.0", "5.6", "8.5"])

add_heading(doc, "异常与降级逻辑", 2)
add_bullets(doc, [
    "在线模型不可用  自动转入离线缓存或规则建议，仍可继续演示。",
    "Agent输出越过白名单或参数范围  拒绝该字段并使用合法默认值或要求用户确认。",
    "回测数据或指标不满足校验  隔离不可信字段并提示，不让Agent把异常数值写入结论。",
    "用户未确认配置  不启动真实回测。",
    "最终方案未由评审委员会确认  不标记为可应用方案，也不执行资源划拨。",
])

add_heading(doc, "页面视觉与动画顺序", 2)
add_bullets(doc, [
    "先出现左侧产品定位、设计原则和可信边界，回答系统为何这样设计。",
    "中部从数据层L1向上讲到交互与治理层L4，讲到L2时停顿并强调确定性计算核心。",
    "随后出现右侧纵向可信治理能力，说明输入留痕、版本、参数、结果和人工决定均可追踪。",
    "最后出现底部业务闭环，把表达偏好、配置确认、回测、贡献证据、双轨建议和评审终裁串起来。",
])

add_heading(doc, "现场讲解词", 2)
add_script_box(doc, "建议35秒", "系统中部是四层能力。底层固定HS300数据、行业映射、因子版本和运行日志；确定性金融计算层完成40项因子输入、组合构建、交易成本、32个联盟反事实回测和精确Shapley；Agent层只把自然语言转成白名单内的结构化建议，并经过参数校验和人工确认；最上层负责证据展示和评审终裁。右侧的可信治理能力贯穿四层，底部形成从偏好输入到评审终裁的业务闭环。收益、风险和角色贡献均由确定性引擎计算，LLM不能生成或改写金融事实。")

add_heading(doc, "页间衔接", 2)
add_table(doc, ["方向", "衔接话术"], [
    ["承接第7页", "刚才的五步证据链之所以能够复算，是因为系统从数据、计算、Agent到治理划分了清晰责任。"],
    ["离开本页", "在这个可信边界下，我们才能进一步讨论贡献结果本身和它对治理的实际价值。"],
    ["进入第13页", "当角色贡献已经可复算，下一步不是让算法直接分配，而是把证据以渐进方式接入机构现行制度。"],
], ["3.2", "13.5"])

add_heading(doc, "可能问答", 2)
add_table(doc, ["评委可能提问", "建议回答"], [
    ["没有API Key还能演示吗", "可以。项目自带经校验的离线Agent缓存，命中预设场景时可完整展示建议、确认和回测；未命中时由规则Agent降级。金融计算始终在本地确定性引擎完成。"],
    ["LLM是否会产生幻觉", "自然语言建议可能有不确定性，所以系统设置因子白名单、参数校验、证据编号和人工确认。模型不能写入金融结果字段。"],
    ["为什么不全部用规则", "规则适合计算和约束，Agent适合理解多样化偏好和组织可读解释。两者分层能同时获得灵活交互与可复算结果。"],
    ["更换模型会不会改变结果", "可能改变配置建议的表达，但经过人工确认后，同一生效配置进入同一引擎会得到同一金融结果。"],
], ["5.0", "11.7"])

# Slide 13
doc.add_page_break()
add_heading(doc, "第13页 6比4双轨治理", 1)
add_body(doc, "页面标题  6比4双轨让贡献证据进入治理，但不替代治理")
add_callout(doc, "页面任务", "说明Shapley证据怎样转化为可落地的分配建议，同时保持制度连续性、角色表达权和评审委员会终裁权。", fill=PALE_AMBER, accent=AMBER)
add_slide_image(doc, "slide-13.png", "第13页内容原型  制度轨与证据轨合成后进入评审委员会决策门")

add_heading(doc, "核心概念", 2)
add_table(doc, ["概念", "含义", "边界"], [
    ["制度轨60%", "沿用机构提交并确认的现行分配份额", "不是系统推断的历史制度"],
    ["证据轨40%", "按正向Shapley贡献形成的动态建议权重", "只占建议总份额的40%，不等于原始Shapley占比"],
    ["带符号Shapley", "保留每个角色对超额收益的正向或负向边际影响", "负值保留用于解释和复盘"],
    ["合成建议", "制度轨和证据轨按固定比例加权的系统建议", "不是最终决定，也不自动执行"],
    ["角色陈述", "五角色基于证据说明贡献来源、约束和改进计划", "不进入份额计算"],
    ["评审委员会终裁", "独立判断、调整并确认后形成可应用方案", "必须保留人工责任与调整理由"],
], ["3.5", "7.5", "8.7"])

add_heading(doc, "合成公式", 2)
add_equation(doc, "建议份额ᵢ = 60% × 现行份额ᵢ + 40% × 证据池份额ᵢ", "对全部角色求和后仍为100%；最终应用前必须经评审委员会确认")
add_table(doc, ["符号", "含义", "数据来源"], [
    ["i", "某一投研角色", "行业、个股、策略、量化或风控"],
    ["现行份额ᵢ", "机构当前制度下分配给角色i的比例", "用户录入或机构档案"],
    ["证据池份额ᵢ", "角色i在证据轨40%内部的建议权重", "正向Shapley经归一化和上限规则处理"],
    ["建议份额ᵢ", "系统展示给评审委员会的合成比例", "按6比4公式计算"],
], ["3.0", "8.3", "8.4"])
add_callout(doc, "比例口径", "6比4是本项目为渐进接入设计的产品参数，不是大赛规定，也不是数学上唯一最优比例。机构落地时可经治理程序调整，但本次演示统一使用6比4。", fill=PALE_CORAL, accent=CORAL)

add_heading(doc, "从Shapley到证据池权重", 2)
add_body(doc, "原始Shapley用于解释贡献，证据池权重用于提出奖励或资源配置建议，两者必须分开。当前实现遵循以下顺序。")
add_table(doc, ["步骤", "规则", "原因"], [
    ["一 保留原始贡献", "记录每个角色的带符号Shapley值φᵢ", "完整呈现正向贡献和负向拖累"],
    ["二 提取正向贡献", "正向分数ᵢ = max(φᵢ, 0)", "负贡献保留在证据层，但不因绝对值较大而获得证据轨奖励"],
    ["三 可选制度调节", "正向分数乘以非负角色调节系数，默认系数为1", "允许机构在证据轨内表达合规的制度偏好，不改变原始Shapley"],
    ["四 水位法归一化", "按分数比例归一化，并执行单角色默认50%上限", "降低证据轨过度集中"],
    ["五 边界降级", "若没有任何正贡献，证据轨对五角色等权；正贡献角色过少时最小幅度放宽上限", "保证证据轨可归一化，同时优先不奖励负贡献"],
], ["2.8", "7.9", "9.0"])
add_callout(doc, "负贡献处理", "负贡献不进入证据轨奖励，但角色仍保留制度轨中的现行份额。负贡献也不能被删除，它是复盘市场阶段、约束成本和角色配置的重要证据。", fill=PALE_BLUE, accent=BLUE)

add_heading(doc, "端到端治理步骤", 2)
add_bullets(doc, [
    "一  读取评审委员会确认的现行份额，检查合计为100%。",
    "二  读取确定性引擎产生的带符号Shapley贡献和证据轨内部权重。",
    "三  按6比4公式计算每个角色的合成建议份额和对应资源量。",
    "四  同屏展示制度轨、证据轨、合成建议、原始贡献和审计信息。",
    "五  五类角色分别说明贡献来源、约束条件和改进计划，陈述仅供参考。",
    "六  评审委员会独立判断，可结合组织战略、合规、职责难度和不可量化投入进行调整。",
    "七  记录调整前建议、调整后结果、调整理由、确认人和时间。",
    "八  仅在人工确认后，方案才可作为预算、资源或激励的应用依据；系统本身不执行划拨。",
])

add_heading(doc, "演示计算示例", 2)
add_body(doc, "以下仅用于解释公式，不代表项目正式回测结果。假设五角色现行份额均为20%，证据池内部权重为50%、30%、20%、0%、0%。")
add_table(doc, ["角色", "现行份额", "证据池权重", "制度轨贡献", "证据轨贡献", "合成建议"], [
    ["行业", "20%", "50%", "12%", "20%", "32%"],
    ["个股", "20%", "30%", "12%", "12%", "24%"],
    ["策略", "20%", "20%", "12%", "8%", "20%"],
    ["量化", "20%", "0%", "12%", "0%", "12%"],
    ["风控", "20%", "0%", "12%", "0%", "12%"],
    ["合计", "100%", "100%", "60%", "40%", "100%"],
], ["2.6", "2.8", "3.2", "3.0", "3.0", "3.0"], font_size=8.8)
add_body(doc, "示例说明  即使量化和风控在证据轨中的建议权重为零，它们仍保留制度轨折算后的12%。这正体现了“双轨”的意义  证据进行动态校正，但不一次性推翻现有制度。")

add_heading(doc, "适用的广义价值分配场景", 2)
add_table(doc, ["场景", "可分配资源", "使用方式"], [
    ["年度预算", "研究经费、外部专家费、调研预算", "证据作为下一周期预算配置参考"],
    ["研究资源", "行业覆盖、研究支持、会议与调研机会", "结合角色贡献和战略重点调整"],
    ["项目机会", "重点组合、投委会议题、专项研究", "将历史贡献作为进入机会池的证据之一"],
    ["数据与算力", "数据订阅、终端席位、GPU和回测额度", "按贡献与未来任务需要综合配置"],
    ["激励权益", "绩效奖金、团队激励或长期激励参考", "只能作为量化证据，须符合机构制度与合规要求"],
], ["3.2", "6.0", "8.3"])

add_heading(doc, "页面视觉与动画顺序", 2)
add_bullets(doc, [
    "左侧先出现60%制度轨和40%证据轨，强调两类输入的来源与作用不同。",
    "两轨汇入中部公式，让评委看到它们如何合成为每个角色的建议份额。",
    "右侧依次出现系统展示、角色陈述参考和评审委员会终裁，形成明确的治理链。",
    "底部最后出现产品截图、角色陈述边界和适用场景，结论落在人工确认后方可应用。",
])

add_heading(doc, "现场讲解词", 2)
add_script_box(doc, "建议25秒", "贡献证据不直接决定分配。我们采用6比4双轨，60%保留机构现行方案，40%由正向Shapley证据动态校正。系统同屏展示制度轨、证据轨和合成建议；五类角色可以基于证据陈述意见，但陈述不进入份额计算。最终由评审委员会独立判断、调整并确认，确认后才能用于预算、研究资源、项目机会以及数据和算力等广义价值分配。")

add_heading(doc, "页间衔接", 2)
add_table(doc, ["方向", "衔接话术"], [
    ["进入本页", "前面证明了角色贡献可计算、计算过程可信，但治理问题不能只交给算法。"],
    ["本页收束", "因此，AlphaSplit的价值不是自动分配，而是用可审计证据改善分配讨论，进而提升参与积极性和协作质量。"],
    ["进入效果或商业页", "当这一机制应用到预算、资源和项目机会后，机构能够把贡献证据沉淀为可持续的治理能力。"],
], ["3.2", "13.5"])

add_heading(doc, "可能问答", 2)
add_table(doc, ["评委可能提问", "建议回答"], [
    ["为什么是6比4", "它是渐进接入的产品设计  60%保留制度连续性，40%让新证据产生实质影响。它不是唯一最优比例，机构可按治理程序调整。"],
    ["为什么负贡献不参与证据轨", "原始负贡献仍完整保留用于解释和复盘，但若按绝对值归一化，会出现拖累越大反而奖励越多的问题，因此奖励权重只取正向贡献。"],
    ["角色陈述会不会影响分配", "陈述只为评审提供上下文，不进入系统建议公式。评审委员会可以基于陈述和其他治理因素人工调整，并必须记录理由。"],
    ["系统是否自动发奖金", "不会。系统只输出证据和合成建议；资源划拨必须由机构治理责任人确认并通过现有业务流程执行。"],
    ["能否用于奖金以外场景", "可以。机制面向广义价值分配，包括预算、研究资源、项目机会、数据与算力额度以及激励权益。"],
], ["5.0", "11.7"])

# Consistency and rehearsal
doc.add_page_break()
add_heading(doc, "三页术语与口径一致性", 1)
add_table(doc, ["术语", "统一写法", "不要写成"], [
    ["项目定位", "投研协作激励与贡献计量平台", "自动奖金分配平台"],
    ["五角色", "行业、个股、策略、量化、风控", "基金经理等未纳入当前动态计量的角色"],
    ["输入规模", "五角色乘以八因子，共40项输入", "40个Agent"],
    ["联盟规模", "五角色的32个联盟全枚举", "五次单角色剔除"],
    ["归因对象", "相对沪深300超额收益的角色边际贡献", "未来收益预测"],
    ["Agent权限", "结构化建议与证据约束解释", "生成收益或决定份额"],
    ["证据轨", "按正向Shapley形成内部建议权重", "直接按Shapley绝对值发放"],
    ["角色陈述", "供评审委员会参考", "由角色集体决定份额"],
    ["评审委员会", "独立判断、调整并确认", "形式化盖章"],
    ["6比4", "渐进接入的产品参数", "大赛规定或唯一最优比例"],
    ["最终目标", "以分配机制激励参与和协同，长期促进投资回报改善", "保证提高投资收益"],
], ["3.0", "8.2", "5.5"], font_size=8.8)

add_heading(doc, "三页逻辑检查", 2)
add_table(doc, ["检查问题", "合格标准"], [
    ["第7页是否说明证据从哪里来", "能清楚说出五角色、40项输入、32联盟、精确Shapley和人工治理"],
    ["第8页是否说明谁负责什么", "能清楚区分数据层、确定性金融引擎、Agent和评审委员会"],
    ["第13页是否说明证据如何进入分配", "能解释制度轨、证据轨、负贡献处理、合成公式和人工终裁"],
    ["三页是否共同支持项目定位", "始终把分配视为激励协同的手段，把长期投资回报改善视为目标而非承诺"],
], ["7.1", "9.6"])

add_heading(doc, "排练清单", 2)
add_bullets(doc, [
    "总时长控制在100秒左右  第7页40秒，第8页35秒，第13页25秒。",
    "第一次出现Shapley时只解释为“所有联盟中的平均边际贡献”，不展开数学推导。",
    "第7页只预告6比4，不提前展开比例理由。",
    "第8页必须完整说出“收益、风险和贡献由确定性引擎计算”。",
    "第13页必须完整说出“角色陈述不进入公式，评审委员会独立确认”。",
    "每页只强调一个结论，不逐字阅读卡片内容。",
    "问答中不把历史回测结果表述为未来收益保证。",
    "官方决赛模板发布后，将三页内容迁移到官方母版，但不改变本说明文档中的口径。",
])

add_heading(doc, "制作交付清单", 2)
add_table(doc, ["交付物", "检查项", "状态"], [
    ["第7页", "五步流程、32联盟、审计记录、Agent边界", "内容已定义"],
    ["第8页", "四层职责、确定性核心、离线缓存、审计脊柱", "内容已定义"],
    ["第13页", "6比4公式、证据池规则、评审决策门、产品截图", "内容已定义"],
    ["讲稿", "40秒、35秒、25秒三个版本", "已写入本说明"],
    ["模板迁移", "待官方决赛PPT模板发布后执行", "待完成"],
], ["3.2", "10.4", "3.1"])

for sec in doc.sections:
    add_header_footer(sec)

doc.core_properties.title = "AlphaSplit 决赛答辩三页内容与讲解说明"
doc.core_properties.subject = "第7页 可审计证据链 第8页 四层系统架构 第13页 6比4双轨治理"
doc.core_properties.author = "AlphaSplit 项目组"
doc.core_properties.keywords = "AlphaSplit, Shapley, 反事实回测, 四层架构, 6比4双轨, 决赛答辩"
doc.save(OUTPUT)
print(OUTPUT)
