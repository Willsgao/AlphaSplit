from pathlib import Path
from zipfile import ZipFile

from docx import Document


path = Path(r"F:\wills\competitions\AlphaSplit\决赛答辩PPT_高玉伟三页_HTML\AlphaSplit_决赛答辩三页内容与讲解说明.docx")

with ZipFile(path) as archive:
    bad_member = archive.testzip()
    assert bad_member is None, f"Corrupt DOCX member: {bad_member}"
    media = [name for name in archive.namelist() if name.startswith("word/media/")]

doc = Document(path)
paragraph_text = [paragraph.text for paragraph in doc.paragraphs]
cell_text = [
    cell.text
    for table in doc.tables
    for row in table.rows
    for cell in row.cells
]
all_text = "\n".join(paragraph_text + cell_text)
headings = {
    style: [p.text for p in doc.paragraphs if p.style.name == style]
    for style in ("Heading 1", "Heading 2", "Heading 3")
}

required = [
    "第7页 可审计证据链",
    "第8页 四层系统架构",
    "第13页 6比4双轨治理",
    "v(S) = 联盟S的组合收益 − 同期沪深300收益",
    "建议份额ᵢ = 60% × 现行份额ᵢ + 40% × 证据池份额ᵢ",
    "角色陈述不进入公式",
    "LLM不能生成或改写金融事实",
]
for item in required:
    assert item in all_text, f"Missing required content: {item}"

for forbidden in (
    r"F:\wills",
    "portfolio-v7-20260827",
    r"\frac",
    r"\sum",
    "五角色协商决定分配",
):
    assert forbidden not in all_text, f"Forbidden content found: {forbidden}"

assert len(media) == 3, f"Expected 3 embedded slide images, got {len(media)}"
assert len(doc.tables) >= 25, f"Expected detailed tables, got {len(doc.tables)}"
assert len(headings["Heading 1"]) >= 6
assert len(headings["Heading 2"]) >= 20
assert len(headings["Heading 3"]) >= 5

print(f"DOCX: {path}")
print(f"Size: {path.stat().st_size} bytes")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")
print(f"Embedded images: {len(media)}")
for style, values in headings.items():
    print(f"{style}: {len(values)}")
print("Validation: PASS")
