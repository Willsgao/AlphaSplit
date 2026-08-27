# -*- coding: utf-8 -*-
"""Generate 高玉伟负责章节初稿 Word（v0.2 · 对齐当前工程）。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(r"F:\wills\competitions\AlphaSplit")
FIG = ROOT / "docs" / "figures"
OUTS = [
    Path(r"F:\wills\competitions\材料\AlphaSplit_技术文档_高玉伟章节初稿_v0.1.docx"),
    ROOT / "docs" / "AlphaSplit_技术文档_高玉伟章节初稿_v0.2.docx",
]

doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)


def set_run_font(run, name="宋体", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, "黑体", 16 if level == 1 else 13, True)
    return p


def p(text, size=12, first=True, center=False, bold=False, color=None):
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if first and not center:
        para.paragraph_format.first_line_indent = Cm(0.74)
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(text)
    set_run_font(run, "宋体", size, bold, color)
    return para


def tip(text):
    return p("【说明】" + text, size=10.5, first=False, color=RGBColor(0x66, 0x66, 0x66))


def caption(text):
    return p(text, size=10.5, first=False, center=True, bold=True)


def add_fig(name, width=5.9):
    path = FIG / name
    if not path.exists():
        tip(f"缺图：{name}")
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path), width=Inches(width))


def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(htext)
        set_run_font(run, "宋体", 10.5, True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, "宋体", 10.5)
    doc.add_paragraph()


# ========== Cover ==========
for _ in range(2):
    doc.add_paragraph()
para = doc.add_paragraph()
para.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(para.add_run("第五届中国研究生金融科技创新大赛"), "黑体", 15, True)
para = doc.add_paragraph()
para.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(para.add_run("技术文档 · 分章节初稿"), "宋体", 12)
doc.add_paragraph()
para = doc.add_paragraph()
para.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(para.add_run("AlphaSplit：多智能体驱动的投研角色贡献计量平台"), "黑体", 14, True)
doc.add_paragraph()
para = doc.add_paragraph()
para.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(para.add_run("高玉伟负责部分：摘要 · 第2章前半 · 第5章 · 第7章供料"), "黑体", 12, True)
for line in [
    "版本：v0.2（对齐当前仓库实现，含 HS300 Shapley / 6:4 双轨 / 五角色协商）",
    "日期：2026年8月",
    "注意：第2.3–2.5、第3/4/6章正文由其他成员主笔；本文不替代全文统稿终版",
]:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(para.add_run(line), "宋体", 10.5)

doc.add_page_break()

h("写作者说明")
p(
    "本分册覆盖技术文档分工中高玉伟主笔/收口内容。叙事口径：直接计量单位是投研五角色（行业/个股/策略/量化/风控分析师）；"
    "因子是角色决策输入的代理（套壳）。主线回测引擎（backtest/query_engine.py）对五角色已实现基于全子集 Pipeline 反事实的 Shapley 归因（2^5=32 联盟）；"
    "价值分配采用 6:4 双轨（60% 现行制度 + 40% 夏普利证据池）及五角色 Agent 协商 + 人工确认。"
    "附录中的三因素合成归因、三人 quant_fund 场景仅作机制演示，不与主线混用。"
)
p("图5-4 系列界面截图已于本地 Demo（http://127.0.0.1:8000）自动截取并嵌入。", first=False)

# ========== 摘要 ==========
h("摘要")
tip("统稿时可根据第6章最终数字替换文中演示区间结果；下列数字来自官方路演包（2024H1 HS300）。")
p(
    "资产管理与量化投研是典型的多角色协作过程：行业研究、个股研究、策略择时、量化信号与风险控制等环节共同影响组合业绩。"
    "传统考核往往依赖固定比例或主观评议，难以回答“各投研角色分别贡献了多少、依据是什么”。"
    "经典业绩归因能够把收益拆到配置、选股或风格等维度，但很少继续映射到组织角色。"
)
p(
    "AlphaSplit 构建“多智能体驱动的投研角色贡献计量平台”：以五类投研角色为计量对象，"
    "在沪深300成分股与公开行情上运行多角色 Pipeline（行业配置→个股因子→量化优化→择时→风控），"
    "通过对 32 个角色子集进行反事实回测计算 Shapley 边际贡献，输出可解释的贡献系数与分配建议。"
    "系统定位是协商锚点与决策支持，采用 6:4 双轨分配（60% 沿用机构现行方案、40% 按夏普利证据动态校正），"
    "经五角色 Agent 协商后形成广义价值分配建议，并由治理责任人确认是否应用。"
)
p(
    "当前工程已形成可本地运行闭环：FastAPI + Web 路演演示（约 22–28 秒分阶段回放）、"
    "HS300 实时/缓存回测、SSE 进度流、可替换 Agent 层（agent_layer/）规划与解读、"
    "6:4 双轨合成表与五角色协商 Mock。"
    "官方演示包示例：2024-01-01~2024-06-30 组合累计收益约 21.50%，夏普约 5.38；"
    "夏普利相对贡献最高为行业分析师（约 58.1%），策略分析师呈轻微负向拖累。"
)

add_fig("fig1-1_pipeline.png", 6.2)
caption("图1-1  解题思路总流程")

doc.add_page_break()

# ========== 第2章 ==========
h("第2章  系统架构与角色—决策输入建模（高玉伟：2.1–2.2）")
tip("2.3 角色口径由董斌主供；2.4–2.5 映射与数据由徐远主笔。")

h("2.1 总体架构图与模块划分", 2)
p(
    "系统按四层组织。对外输出与考核口径的对象是角色；因子处于信号层，表达各角色决策输入的代理。"
    "计量层以 HS300 回测引擎为核心，对五角色执行 Shapley 子集反事实；"
    "智能层含可独立替换的 agent_layer（规划/解读）与协商模块；"
    "输出层提供贡献可视化、6:4 双轨合成与人工确认。"
)
add_fig("fig2-1_architecture.png", 6.0)
caption("图2-1  AlphaSplit 系统总体架构")

p("四层含义：", first=False)
p("（1）信号层：data/processed/ 下 HS300 日线、行业映射、财务因子等公开数据。", first=False)
p("（2）计量层：backtest/query_engine.py 定义 v(S) 为「仅启用子集 S 中角色决策输入」时的 Pipeline 组合收益，并计算 Shapley 值。", first=False)
p("（3）智能层：agent_layer/（可替换）负责 NL→回测计划与结果解读；engine/agentic/role_negotiation.py 实现五角色协商。", first=False)
p("（4）输出层：Web 演示、6:4 合成、带宽内建议案与人工确认。", first=False)

p(
    "仓库关键模块：engine/api（FastAPI）、engine/web（路演前端）、engine/integrations/quantskills_backtest.py（桥接）、"
    "backtest/（回测归因引擎）、agent_layer/（独立 Agent 包）、engine/agentic/（协商与附录 Mock）、"
    "data/scenarios/（场景 JSON）。启动：python -m engine.api.main。"
)
add_fig("fig2-2_modules.png", 6.0)
caption("图2-2  仓库模块与调用关系")

p("主数据流（主线）：", first=False)
p(
    "用户/Agent 规划 query → POST /api/backtest/run(/stream) → 五角色因子选股 + Pipeline 回测 "
    "→ Shapley 32 子集反事实 → role_breakdown → 6:4 双轨合成 → 五角色协商 → 人工确认。",
    first=False,
)

h("2.2 技术栈、依赖与运行环境", 2)
table(
    ["类别", "组件", "用途"],
    [
        ["语言", "Python 3.10+", "主业务"],
        ["API", "FastAPI / Uvicorn", "/api/* + 静态页"],
        ["回测引擎", "backtest/query_engine.py", "HS300 五角色 Pipeline + Shapley"],
        ["前端", "HTML/CSS/JS + Chart.js", "路演演示、过程看板、双轨/协商"],
        ["Agent", "agent_layer/（默认 Mock，可选 DeepSeek）", "规划 query、解读报告"],
        ["数据", "data/processed/*.csv", "HS300 成分股行情与行业映射"],
        ["缓存", "backtest/results/_cache", "相同参数秒开（路演）"],
    ],
)
p(
    "可运行性：演示包自带 data/processed/ 与 engine/web/demo/backtest_demo.json；"
    "默认不依赖外网 LLM；一键启动后可完成「路演演示 → 人工确认 → 五角色协商」全流程。"
)
add_fig("fig4-3_mock.png", 5.6)
caption("图4-3  Agent Mock / 可选 LLM（可运行性保障）")

add_fig("fig5-4a_backtest_home.png", 6.0)
caption("图2-3  Web 演示首页（HS300 岗位贡献 + Agent 规划面板）")

h("2.3–2.5 衔接说明", 2)
p(
    "2.3–2.5 由董斌/徐远主笔。统稿时需保证：角色集合与第5章输出维度一致（五分析师）；"
    "并明确因子仅为决策输入代理。行业筛选、role_weights 等参数见回测 HTTP 接口文档。"
)

doc.add_page_break()

# ========== 第5章 ==========
h("第5章  Shapley 贡献计量、双轨分配与人工确认")
tip("主笔：高玉伟；5.3 请徐远按 query_engine 实现核对公式与特征函数定义。")

h("5.1 从边际贡献到 Shapley 值", 2)
p(
    "设角色集合 N={行业,个股,策略,量化,风控}，|N|=5。"
    "对子集 S⊆N，特征函数 v(S) 表示仅保留 S 中角色对应决策模块时的 Pipeline 组合收益（反事实回测）。"
    "角色 i 的 Shapley 值 φ_i 为各加入顺序边际贡献的加权平均，满足有效性、对称性、哑元性等公理。"
)
p("φ_i = Σ_{S⊆N\\{i}} [|S|!(n-|S|-1)! / n!] · [v(S∪{i}) - v(S)]", first=False, center=True)
p(
    "工程实现（query_engine._compute_shapley_attribution）：对五角色遍历 2^5=32 个联盟，"
    "每个联盟执行真实 Pipeline 回测得到 v(S)，再汇总 φ_i。"
    "输出 role_breakdown 含 shapley_raw（带符号边际贡献）、weight（新池 40% 内建议权重）等字段。"
)
add_fig("fig5-1_shapley.png", 5.9)
caption("图5-1  从角色边际贡献到 Shapley 贡献系数")

h("5.2 计算复杂度与工程优化", 2)
p(
    "五角色下全子集 Shapley 需 32 次 Pipeline 级反事实回测，演示区间（如 2024H1）单次约 1–7 分钟；"
    "相同参数命中 backtest/results/_cache 可秒开。"
    "角色数增加时可扩展 Monte Carlo 排列采样或 Truncated Shapley；"
    "答辩应报告：子集数、耗时、缓存命中率及贡献排序稳定性。"
)

h("5.3 与 LOO / 附录归因的区分", 2)
p(
    "Leave-One-Out 仅需 O(n) 次回测，适合敏感性分析，但不是主线默认输出。"
    "engine/core/backtest.py 与 POST /api/attribution 提供的「策略/行业·政策/偶然·宏观」三因素对照，"
    "用于解释超额质量结构，与 HS300 五角色 Shapley 主线互补，答辩与文档中应分开展示。"
)
add_fig("fig5-3_loo_shapley.png", 5.9)
caption("图5-3  LOO 与全子集 Shapley 的关系（方法对照）")

h("5.4 6:4 双轨分配机制", 2)
p(
    "产品采用双轨制（详见 docs/双轨分配机制.md）：可分配资源池 P 中，"
    "60% 沿用机构提交的现行分配方案（旧池），40% 按回测夏普利证据池建议权重形成方案。"
    "各角色建议份额：建议_i = 60%×现行_i + 40%×夏普利证据池_i。"
    "该设计在立项「协商锚点 + 人工确认」框架下，把证据渐进接入制度而非全盘推翻历史方案。"
)
table(
    ["轨道", "比例", "含义", "数据来源"],
    [
        ["旧池", "60%", "现行制度/治理责任人确认比例", "用户录入（演示默认五角色各 20%）"],
        ["证据池", "40%", "夏普利证据建议", "role_breakdown.weight"],
        ["合成", "100%", "考核讨论基准", "Web 双轨合成表"],
    ],
)
add_fig("fig5-4c_dual_pool.png", 6.0)
caption("图5-4  6:4 双轨分配合成表（人工确认页）")

h("5.5 五角色 Agent 协商与可替换 Agent 层", 2)
p(
    "人工确认页主线为五角色 Agent 协商（engine/agentic/role_negotiation.py）："
    "五岗位各基于夏普利证据与 6:4 合成锚点陈述主张，投委会仲裁后形成带宽内建议案。"
    "与早期 quant_fund 三人因子复用场景不同，协商对象与回测 role_breakdown 一一对应。"
)
p(
    "回测规划与解读由独立目录 agent_layer/ 实现（经 engine/integrations/agent_bridge.py 桥接），"
    "队员可通过 agent_layer/registry.py 或环境变量 AGENT_LAYER_IMPL 整体替换实现，"
    "默认 Mock 规则解析自然语言（如「2024上半年半导体、偏重行业岗」）为 query JSON。"
)
add_fig("fig5-4e_agent_panel.png", 6.0)
caption("图5-5  可替换 Agent 层：规划回测面板")

add_fig("fig5-4d_five_role_negotiate.png", 6.0)
caption("图5-6  五角色 Agent 协商与建议案")

h("5.6 关键接口与演示结果", 2)
table(
    ["接口", "作用"],
    [
        ["POST /api/backtest/run", "HS300 五角色回测 + Shapley"],
        ["POST /api/backtest/run/stream", "SSE 流式进度（子集/阶段）"],
        ["GET /api/backtest/demo", "路演官方预置结果（秒开）"],
        ["POST /api/agent/plan", "Agent 规划 query（agent_layer）"],
        ["POST /api/agent/interpret", "Agent 解读回测报告"],
        ["POST /api/negotiate", "五角色协商（传 backtest + baseline_weights）"],
        ["POST /api/confirm", "人工确认分配方案（role_model=five_roles）"],
    ],
)
p(
    "演示包（2024H1）结果摘要：组合累计收益约 21.50%；夏普利相对贡献排序为"
    "行业分析师（58.1%）> 量化分析师（16.4%）> 风控分析师（13.8%）> 个股分析师（13.1%）> 策略分析师（-1.4%）；"
    "新池 40% 内行业分析师建议权重约 53.5%。",
    first=False,
)
add_fig("fig5-4b_backtest_conclusion.png", 6.0)
caption("图5-7  回测结论摘要与夏普利排名（路演演示）")

doc.add_page_break()

# ========== 第7章供料 ==========
h("第7章供料（给汪皓怡主笔用）")

h("7.1 里程碑素材", 2)
table(
    ["阶段", "事项", "产出"],
    [
        ["立项", "贡献计量主线；突出广义价值分配", "课题立项描述"],
        ["工程骨架", "FastAPI + Web Demo", "可本地启动"],
        ["回测引擎接入", "HS300 + Shapley 32 子集", "query_engine + 桥接层"],
        ["路演前端", "分阶段演示、底部进度条、过程看板", "engine/web"],
        ["双轨分配", "6:4 合成 + 文档", "双轨分配机制.md"],
        ["Agent 独立包", "agent_layer 可替换", "规划/解读 API"],
        ["五角色协商", "对齐夏普利证据", "role_negotiation.py"],
        ["文档 v0.2", "高玉伟章节 + Demo 截图", "本分册"],
    ],
)

h("7.2 关键取舍", 2)
p("（1）计量对象=五投研角色，因子仅作套壳代理。", first=False)
p("（2）主线用全子集 Shapley（n=5 可承受），非仅 LOO。", first=False)
p("（3）分配用 6:4 双轨，尊重现行制度 + 证据校正。", first=False)
p("（4）Agent 默认 Mock；LLM 可选 DeepSeek 兼容接口。", first=False)
p("（5）三人 quant_fund 场景降为附录，避免与五角色主线混淆。", first=False)

h("7.3 局限与改进", 2)
p("套壳与真实岗位决策差距；子集反事实假设主观性；五角色下计算耗时；机构脱敏数据缺乏；行业全板块并行回测属扩展项。", first=False)
p("改进：脱敏决策日志、MC Shapley 扩展、Agent 执行重试、第6章完整量化表、一键复现脚本。", first=False)

doc.add_page_break()
h("附录：本分册插图清单")
table(
    ["图号", "文件", "用途"],
    [
        ["图1-1", "fig1-1_pipeline.png", "解题总流程"],
        ["图2-1", "fig2-1_architecture.png", "总体架构"],
        ["图2-2", "fig2-2_modules.png", "模块关系"],
        ["图2-3", "fig5-4a_backtest_home.png", "Web 首页"],
        ["图4-3", "fig4-3_mock.png", "Mock/LLM"],
        ["图5-1", "fig5-1_shapley.png", "Shapley 示意"],
        ["图5-3", "fig5-3_loo_shapley.png", "LOO 对照"],
        ["图5-4", "fig5-4c_dual_pool.png", "6:4 双轨表"],
        ["图5-5", "fig5-4e_agent_panel.png", "Agent 规划"],
        ["图5-6", "fig5-4d_five_role_negotiate.png", "五角色协商"],
        ["图5-7", "fig5-4b_backtest_conclusion.png", "回测结论"],
    ],
)

for o in OUTS:
    o.parent.mkdir(parents=True, exist_ok=True)
    doc.save(o)
    print("saved", o)
