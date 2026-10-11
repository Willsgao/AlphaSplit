from __future__ import annotations

import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "AlphaSplit_完整技术文档_初赛提交版_v1.3.docx"
ASSET_DIR = ROOT / "docs" / "_full_tech_doc_assets_v1"
DEMO_JSON = ROOT / "engine" / "web" / "demo" / "backtest_demo.json"

FONT_CJK = Path(r"C:\Windows\Fonts\NotoSansSC-VF.ttf")
FONT_CJK_NAME = "Noto Sans SC"
FONT_LATIN = "Calibri"

NAVY = "0B2545"
BLUE = "2457A6"
CYAN = "16B8C8"
GOLD = "D89B24"
PURPLE = "7457C8"
GREEN = "2E8B72"
RED = "C94C4C"
INK = "18212B"
MUTED = "5F6B78"
LIGHT = "F3F7FA"
LIGHT_BLUE = "E8F2F7"
LIGHT_CYAN = "E6F7F8"
LIGHT_GOLD = "FFF6DF"
WHITE = "FFFFFF"
GRID = "D7E1E8"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    if sum(widths_dxa) != 9360:
        raise ValueError(f"table widths must sum to 9360 DXA: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def font_run(run, size=None, bold=None, color=INK, italic=None, latin=FONT_LATIN, east=FONT_CJK_NAME):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)
    return run


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK_NAME)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 8),
        ("Subtitle", 14, MUTED, 0, 10),
        ("Heading 1", 18, NAVY, 16, 8),
        ("Heading 2", 14, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        st = doc.styles[name]
        st.font.name = FONT_LATIN
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK_NAME)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.08

    caption = doc.styles["Caption"]
    caption.font.name = FONT_LATIN
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK_NAME)
    caption.font.size = Pt(9)
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    font_run(run, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run_el = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = "1"
    run_el.append(text_el)
    fld.append(run_el)
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    font_run(run, size=8.5, color=MUTED)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    font_run(p.add_run("AlphaSplit  |  投研贡献计量与广义价值分配"), size=8.5, bold=True, color=MUTED)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), GRID)
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_field(fp)
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""


def add_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    base_abs = max(abstract_ids or [0]) + 1
    base_num = max(num_ids or [0]) + 1

    def create(abstract_id: int, num_id: int, fmt: str, text: str, font: str | None = None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(abstract_id))
        num.append(abs_id)
        numbering.append(num)

    create(base_abs, base_num, "bullet", "•", FONT_LATIN)
    create(base_abs + 1, base_num + 1, "decimal", "%1.")
    return base_num, base_num + 1


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def add_bullet(doc, text: str, bullet_num_id: int) -> None:
    p = doc.add_paragraph()
    apply_num(p, bullet_num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    font_run(p.add_run(text), size=10.5, color=INK)


def add_numbered(doc, text: str, decimal_num_id: int) -> None:
    p = doc.add_paragraph()
    apply_num(p, decimal_num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    font_run(p.add_run(text), size=10.5, color=INK)


def add_body(doc, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = False
    if bold_prefix and text.startswith(bold_prefix):
        font_run(p.add_run(bold_prefix), size=10.5, bold=True, color=NAVY)
        font_run(p.add_run(text[len(bold_prefix):]), size=10.5, color=INK)
    else:
        font_run(p.add_run(text), size=10.5, color=INK)


def add_callout(doc, label: str, text: str, fill=LIGHT_CYAN, accent=CYAN) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    font_run(p.add_run(f"{label}  "), size=10.5, bold=True, color=NAVY)
    font_run(p.add_run(text), size=10.5, color=INK)


def add_equation(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    font_run(p.add_run(text), size=11, bold=True, color=NAVY, latin="Cambria Math")


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], small=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        font_run(p.add_run(text), size=8.6 if small else 9, bold=True, color=WHITE)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2 == 1:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            font_run(p.add_run(str(value)), size=8.3 if small else 8.8, color=INK)
    set_table_geometry(table, widths)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, path: Path, caption: str, width=6.35, alt_text: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    if alt_text:
        inline._inline.docPr.set("descr", alt_text)
        inline._inline.docPr.set("title", caption)
    cp = doc.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(cp.add_run(caption), size=9, color=MUTED)


def chapter_start(doc, chapter: str, title: str, lead: str) -> None:
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    font_run(p.add_run(chapter.upper()), size=9, bold=True, color=CYAN)
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_before = Pt(0)
    font_run(h.add_run(title), size=21, bold=True, color=NAVY)
    add_callout(doc, "本章要点", lead, fill=LIGHT_BLUE, accent=BLUE)


FONT_BOLD_CJK = Path(r"C:\Windows\Fonts\simhei.ttf")


def img_font(size: int, bold: bool = False):
    return ImageFont.truetype(str(FONT_BOLD_CJK if bold else FONT_CJK), size=size)


def text_width(draw, text: str, font) -> int:
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0]


def wrap_pixels(draw, text: str, font, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in str(text).split("\n"):
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for char in paragraph:
            candidate = current + char
            if current and text_width(draw, candidate, font) > max_width:
                lines.append(current)
                current = char
            else:
                current = candidate
        if current:
            lines.append(current)
    return lines


def centered_text(draw, box, text: str, font, fill, spacing=8):
    x0, y0, x1, y1 = box
    lines = wrap_pixels(draw, text, font, x1 - x0 - 20)
    heights = [draw.textbbox((0, 0), line or " ", font=font)[3] for line in lines]
    total = sum(heights) + spacing * max(0, len(lines) - 1)
    y = y0 + (y1 - y0 - total) / 2
    for line, h in zip(lines, heights):
        w = text_width(draw, line, font)
        draw.text((x0 + (x1 - x0 - w) / 2, y), line, font=font, fill=fill)
        y += h + spacing


def draw_arrow(draw, start, end, color="#93A4B2", width=5):
    x0, y0 = start
    x1, y1 = end
    draw.line((x0, y0, x1, y1), fill=color, width=width)
    angle = math.atan2(y1 - y0, x1 - x0)
    size = 17
    left = (x1 - size * math.cos(angle - 0.55), y1 - size * math.sin(angle - 0.55))
    right = (x1 - size * math.cos(angle + 0.55), y1 - size * math.sin(angle + 0.55))
    draw.polygon([end, left, right], fill=color)


def draw_dashed_polyline(draw, points, color, width=5, dash=18, gap=10):
    for (x0, y0), (x1, y1) in zip(points, points[1:]):
        length = math.hypot(x1 - x0, y1 - y0)
        if length == 0:
            continue
        ux, uy = (x1 - x0) / length, (y1 - y0) / length
        cursor = 0.0
        while cursor < length:
            end = min(cursor + dash, length)
            draw.line((x0 + ux * cursor, y0 + uy * cursor,
                       x0 + ux * end, y0 + uy * end), fill=color, width=width)
            cursor += dash + gap


def save_canvas(image: Image.Image, path: Path):
    image.save(path, format="PNG", optimize=True)


def draw_flow(path: Path, title: str, steps: list[tuple[str, str]], colors: list[str] | None = None):
    colors = colors or [BLUE, CYAN, GREEN, GOLD, PURPLE, RED]
    w, h = 2200, 700
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 48), title, font=img_font(46, True), fill=f"#{NAVY}")
    n = len(steps)
    gap = 42
    left, right = 70, 70
    box_w = int((w - left - right - gap * (n - 1)) / n)
    top, box_h = 190, 350
    for i, (head, body) in enumerate(steps):
        x0 = left + i * (box_w + gap)
        x1 = x0 + box_w
        color = f"#{colors[i % len(colors)]}"
        d.rounded_rectangle((x0, top, x1, top + box_h), radius=28, fill="#F7FAFC", outline=color, width=4)
        d.text((x0 + 26, top + 23), f"{i + 1:02d}", font=img_font(26, True), fill=color)
        centered_text(d, (x0 + 16, top + 70, x1 - 16, top + 185), head, img_font(31, True), f"#{NAVY}", 4)
        centered_text(d, (x0 + 20, top + 180, x1 - 20, top + 315), body, img_font(23), f"#{MUTED}", 8)
        if i < n - 1:
            draw_arrow(d, (x1 + 7, top + box_h // 2), (x1 + gap - 8, top + box_h // 2), width=4)
    d.text((70, 610), "所有金融数值由确定性引擎计算，Agent只生成结构化建议。", font=img_font(24), fill=f"#{MUTED}")
    save_canvas(img, path)


def draw_architecture(path: Path):
    layers = [
        ("交互与治理层", "四步自助流程 · 数据依据 · 三方对照 · 角色陈述参考 · 评审决策", CYAN),
        ("Agent 编排层", "DeepSeek结构化建议 · 规则降级 · 查询规划 · 证据约束解读", PURPLE),
        ("金融计算层", "40因子评分 · 月度调仓 · 风控投影 · 交易成本 · 32联盟Shapley", BLUE),
        ("数据与审计层", "HS300行情 · 行业映射 · 120日预热 · 缓存版本 · 月度诊断", GREEN),
    ]
    w, h = 2100, 1050
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((90, 55), "AlphaSplit 分层架构", font=img_font(50, True), fill=f"#{NAVY}")
    for i, (name, body, color) in enumerate(layers):
        y0 = 190 + i * 190
        y1 = y0 + 135
        d.rounded_rectangle((125, y0, 1975, y1), radius=28, fill="#F7FAFC", outline=f"#{color}", width=4)
        d.rounded_rectangle((125, y0, 520, y1), radius=28, fill=f"#{color}", outline=f"#{color}")
        centered_text(d, (135, y0 + 5, 510, y1 - 5), name, img_font(30, True), "white")
        centered_text(d, (560, y0 + 5, 1920, y1 - 5), body, img_font(27), f"#{INK}")
        if i < len(layers) - 1:
            draw_arrow(d, (1050, y1 + 10), (1050, y0 + 178), width=4)
    centered_text(d, (200, 930, 1900, 1010), "LLM只负责语义与结构化建议；收益、风险和贡献全部由确定性引擎计算",
                  img_font(25), f"#{MUTED}")
    save_canvas(img, path)


def draw_role_map(path: Path):
    roles = [
        ("行业分析师", "行业景气与赛道\n4行业规模中性聚合", "动量 · 趋势 · 突破 · 位置", BLUE),
        ("个股分析师", "候选股排序\n保留30只候选", "反转 · RSI · 形态 · 回撤", CYAN),
        ("策略分析师", "择时与仓位\n弱趋势防御", "收益 · 波动 · ATR · 分布", GOLD),
        ("量化分析师", "联合排序优化\n与个股角色50/50", "量能 · 流动性 · 量价 · 动量", PURPLE),
        ("风控分析师", "联合容量投影\n单股10% / 行业30%", "波动 · 下行 · 回撤 · 尾部", RED),
    ]
    w, h = 2200, 980
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 45), "五角色 × 40因子：从岗位判断到可计算输入", font=img_font(47, True), fill=f"#{NAVY}")
    gap, left = 34, 70
    box_w = int((w - 140 - gap * 4) / 5)
    top, bottom = 180, 810
    for i, (role, action, factors, color) in enumerate(roles):
        x0 = left + i * (box_w + gap)
        x1 = x0 + box_w
        d.rounded_rectangle((x0, top, x1, bottom), radius=30, fill="#F8FBFD", outline=f"#{color}", width=4)
        d.rounded_rectangle((x0, top, x1, top + 135), radius=30, fill=f"#{color}", outline=f"#{color}")
        centered_text(d, (x0 + 8, top + 8, x1 - 8, top + 128), role, img_font(29, True), "white")
        centered_text(d, (x0 + 20, top + 175, x1 - 20, top + 360), action, img_font(25), f"#{INK}", 11)
        d.line((x0 + 35, top + 390, x1 - 35, top + 390), fill="#D7E1E8", width=3)
        centered_text(d, (x0 + 20, top + 420, x1 - 20, top + 545), factors, img_font(22), f"#{MUTED}", 8)
        centered_text(d, (x0 + 12, bottom - 75, x1 - 12, bottom - 22), "8个因子 · 角色内归一化", img_font(20, True), f"#{color}")
    centered_text(d, (100, 860, 2100, 950), "默认每个角色内部8因子等权12.5%，仅是透明初始化基线；Agent可在白名单和边界内生成局部补丁",
                  img_font(24), f"#{MUTED}")
    save_canvas(img, path)


def draw_governance(path: Path):
    w, h = 2100, 850
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 55), "6:4 双轨治理：贡献证据渐进进入资源配置", font=img_font(47, True), fill=f"#{NAVY}")
    boxes = [
        (110, 230, 590, 575, "制度轨 60%", "沿用机构现行方案\n保留治理连续性", BLUE),
        (755, 230, 1235, 575, "证据轨 40%", "正向Shapley形成建议权重\n单角色上限50%", CYAN),
        (1400, 230, 1990, 575, "评审建议 100%", "角色陈述仅供参考\n评审委员会最终决策", GREEN),
    ]
    for x0, y0, x1, y1, title, body, color in boxes:
        d.rounded_rectangle((x0, y0, x1, y1), radius=32, fill="#F7FAFC", outline=f"#{color}", width=5)
        centered_text(d, (x0 + 20, y0 + 35, x1 - 20, y0 + 155), title, img_font(35, True), f"#{color}")
        centered_text(d, (x0 + 30, y0 + 155, x1 - 30, y1 - 35), body, img_font(27), f"#{INK}", 12)
    draw_arrow(d, (610, 402), (735, 402), width=5)
    centered_text(d, (620, 305, 725, 375), "+", img_font(38, True), f"#{MUTED}")
    draw_arrow(d, (1255, 402), (1380, 402), width=5)
    centered_text(d, (180, 645, 1920, 715), "最终比例 = 60% × 制度轨内部比例 + 40% × 证据轨内部比例",
                  img_font(28, True), f"#{NAVY}")
    centered_text(d, (100, 735, 2000, 820), "系统输出支持预算、研究资源、项目机会、数据与算力额度等广义价值分配，不自动执行资源划拨",
                  img_font(23), f"#{MUTED}")
    save_canvas(img, path)


def draw_axes(draw, rect, y_min, y_max, y_ticks, x_labels, x_positions, y_suffix="%"):
    x0, y0, x1, y1 = rect
    draw.line((x0, y1, x1, y1), fill="#AAB7C2", width=3)
    for value in y_ticks:
        y = y1 - (value - y_min) / (y_max - y_min) * (y1 - y0)
        draw.line((x0, y, x1, y), fill="#E6EDF2", width=2)
        label = f"{value:g}{y_suffix}"
        draw.text((x0 - 78, y - 13), label, font=img_font(19), fill=f"#{MUTED}")
    for pos, label in zip(x_positions, x_labels):
        draw.text((pos - text_width(draw, label, img_font(18)) / 2, y1 + 16), label, font=img_font(18), fill=f"#{MUTED}")


def make_figures(demo: dict) -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    paths = {name: ASSET_DIR / name for name in (
        "fig01_solution.png", "fig02_architecture.png", "fig03_role_map.png",
        "fig04_pipeline.png", "fig05_shapley.png", "fig06_agent.png",
        "fig07_governance.png", "fig08_returns.png", "fig09_roles.png",
        "fig10_audit.png", "fig11_robustness.png", "fig12_validation.png",
    )}

    draw_flow(paths["fig01_solution.png"], "从客户意图到可审计贡献证据", [
        ("描述需求", "快捷选择 +\n自然语言"),
        ("确认方案", "Agent生成白名单\n因子参数补丁"),
        ("真实回测", "优化方案 / 基线 /\nHS300同口径"),
        ("贡献计量", "32联盟反事实\n精确Shapley"),
        ("评审决策", "6:4建议 +\n角色陈述参考"),
    ])
    draw_architecture(paths["fig02_architecture.png"])
    draw_role_map(paths["fig03_role_map.png"])
    draw_flow(paths["fig04_pipeline.png"], "月度组合构建与收益计算流水线", [
        ("120日预热", "只计算信号\n不计入收益"),
        ("行业筛选", "Top-4行业\n规模中性Borda"),
        ("联合选股", "个股/量化\n50%/50%排名"),
        ("仓位与风控", "趋势仓位 +\n联合容量投影"),
        ("次月持有", "10只股票\n按日收益复利"),
        ("成本审计", "实际单边换手 ×\n费率与滑点"),
    ])
    draw_flow(paths["fig05_shapley.png"], "精确Shapley反事实计量", [
        ("角色集合", "5个角色\nN={I,S,T,Q,R}"),
        ("联盟枚举", "31个非空联盟\n+ 1个空联盟"),
        ("统一回测", "同数据/区间/成本\n仅角色集合变化"),
        ("边际贡献", "比较v(S∪{i})\n与v(S)"),
        ("加权汇总", "所有加入顺序\n平均边际贡献"),
        ("有效性校验", "Σφᵢ =\n组合超额收益"),
    ], [BLUE, CYAN, GREEN, GOLD, PURPLE, RED])
    draw_flow(paths["fig06_agent.png"], "受约束Agent配置闭环", [
        ("用户表达", "风险偏好/目标/\n关注角色"),
        ("DeepSeek理解", "输出结构化\n配置建议"),
        ("确定性校验", "角色/因子白名单\n范围与归一化"),
        ("变更预览", "当前值→建议值\n理由与审计号"),
        ("人工确认", "未确认不生效\n可恢复默认"),
        ("回测与解释", "确定性计算 +\n证据编号引用"),
    ], [CYAN, PURPLE, BLUE, GOLD, GREEN, RED])
    draw_governance(paths["fig07_governance.png"])

    comparison = demo["comparison"]
    periods = [x["period"] for x in comparison["customer_time_series"]]
    def cumulative(items):
        value = 1.0
        out = [0.0]
        for item in items:
            value *= 1 + float(item["return"])
            out.append((value - 1) * 100)
        return out
    xlabels = ["起点"] + periods
    customer = cumulative(comparison["customer_time_series"])
    default = cumulative(comparison["default_time_series"])
    benchmark_value = 1.0
    benchmark = [0.0]
    for item in comparison["customer_time_series"]:
        benchmark_value *= 1 + float(item["benchmark"])
        benchmark.append((benchmark_value - 1) * 100)
    w, h = 2200, 1050
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((90, 45), "同区间三方案累计收益路径（2024-07 至 2025-06）", font=img_font(46, True), fill=f"#{NAVY}")
    plot = (190, 220, 2070, 860)
    y_min, y_max = -6, 27
    x_pos = [plot[0] + i * (plot[2] - plot[0]) / (len(xlabels) - 1) for i in range(len(xlabels))]
    tick_idx = list(range(0, len(xlabels), 2))
    if tick_idx[-1] != len(xlabels) - 1:
        tick_idx.append(len(xlabels) - 1)
    draw_axes(d, plot, y_min, y_max, [-5, 0, 5, 10, 15, 20, 25], [xlabels[i] for i in tick_idx], [x_pos[i] for i in tick_idx])
    def ymap(v):
        return plot[3] - (v - y_min) / (y_max - y_min) * (plot[3] - plot[1])
    for series, color, label, dashed in (
        (customer, CYAN, "优化策略", False),
        (default, GOLD, "严格旧参数基线", False),
        (benchmark, PURPLE, "沪深300", True),
    ):
        pts = [(x_pos[i], ymap(v)) for i, v in enumerate(series)]
        if dashed:
            draw_dashed_polyline(d, pts, f"#{color}", width=5)
        else:
            d.line(pts, fill=f"#{color}", width=6, joint="curve")
        for x, y in pts:
            d.ellipse((x - 7, y - 7, x + 7, y + 7), fill="white", outline=f"#{color}", width=4)
    legend_y = 135
    legend_x = 880
    for label, color in (("优化策略", CYAN), ("严格旧参数基线", GOLD), ("沪深300", PURPLE)):
        d.line((legend_x, legend_y + 15, legend_x + 45, legend_y + 15), fill=f"#{color}", width=6)
        d.text((legend_x + 58, legend_y), label, font=img_font(23), fill=f"#{INK}")
        legend_x += 350
    for value, color, offset in ((customer[-1], CYAN, -35), (default[-1], GOLD, 20), (benchmark[-1], PURPLE, -28)):
        d.text((plot[2] - 110, ymap(value) + offset), f"{value:.2f}%", font=img_font(22, True), fill=f"#{color}")
    save_canvas(img, paths["fig08_returns.png"])

    roles = demo["results"]["role_breakdown"]
    role_names = [x["role"].replace("分析师", "") for x in roles]
    contrib = [x["shapley_raw"] * 100 for x in roles]
    weights = [x["weight"] * 100 for x in roles]
    role_colors = [f"#{BLUE}", f"#{CYAN}", f"#{GOLD}", f"#{PURPLE}", f"#{RED}"]
    w, h = 2200, 920
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 40), "角色贡献与证据轨权重并非同一概念", font=img_font(45, True), fill=f"#{NAVY}")
    panels = [(120, 180, 1040, 790), (1180, 180, 2100, 790)]
    for panel, title in zip(panels, ("Shapley超额贡献（带符号）", "证据轨建议权重（仅正贡献）")):
        d.text((panel[0], 125), title, font=img_font(29, True), fill=f"#{NAVY}")
    # contribution panel
    p = panels[0]
    zero_y = p[1] + (22 / 32) * (p[3] - p[1])
    d.line((p[0], zero_y, p[2], zero_y), fill="#9BAAB5", width=3)
    bar_w = 115
    for i, (name, value, color) in enumerate(zip(role_names, contrib, role_colors)):
        cx = p[0] + 105 + i * 170
        y = zero_y - value / 32 * (p[3] - p[1])
        d.rectangle((cx - bar_w / 2, min(y, zero_y), cx + bar_w / 2, max(y, zero_y)), fill=color)
        d.text((cx - text_width(d, name, img_font(22)) / 2, p[3] + 18), name, font=img_font(22), fill=f"#{MUTED}")
        label = f"{value:+.2f}%"
        ly = y - 34 if value >= 0 else y + 8
        d.text((cx - text_width(d, label, img_font(20, True)) / 2, ly), label, font=img_font(20, True), fill=color)
    # weight panel
    p = panels[1]
    base_y = p[3]
    d.line((p[0], base_y, p[2], base_y), fill="#9BAAB5", width=3)
    for i, (name, value, color) in enumerate(zip(role_names, weights, role_colors)):
        cx = p[0] + 105 + i * 170
        top_y = base_y - value / 55 * (p[3] - p[1])
        d.rectangle((cx - bar_w / 2, top_y, cx + bar_w / 2, base_y), fill=color)
        d.text((cx - text_width(d, name, img_font(22)) / 2, p[3] + 18), name, font=img_font(22), fill=f"#{MUTED}")
        label = f"{value:.0f}%"
        d.text((cx - text_width(d, label, img_font(20, True)) / 2, top_y - 34), label, font=img_font(20, True), fill=color)
    save_canvas(img, paths["fig09_roles.png"])

    diagnostics = demo["results"]["portfolio_diagnostics"]
    months = [x["month"] for x in diagnostics]
    exposure = [x["gross_exposure"] * 100 for x in diagnostics]
    timing = [x["position_ratio"] * 100 for x in diagnostics]
    turnover = [x["turnover"] * 100 for x in diagnostics]
    cost = [x["transaction_cost"] * 100 for x in diagnostics]
    w, h = 2200, 1180
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 35), "逐月组合审计：仓位、换手与交易成本", font=img_font(45, True), fill=f"#{NAVY}")
    left, right = 180, 2070
    x_pos = [left + i * (right - left) / (len(months) - 1) for i in range(len(months))]
    p1 = (left, 170, right, 530)
    draw_axes(d, p1, 0, 110, [0, 25, 50, 75, 100], [], [])
    def map1(v): return p1[3] - v / 110 * (p1[3] - p1[1])
    for series, color in ((exposure, CYAN), (timing, BLUE)):
        pts = [(x_pos[i], map1(v)) for i, v in enumerate(series)]
        d.line(pts, fill=f"#{color}", width=5)
        for x, y in pts:
            d.ellipse((x - 6, y - 6, x + 6, y + 6), fill="white", outline=f"#{color}", width=3)
    d.text((1350, 115), "— 实际股票仓位", font=img_font(22), fill=f"#{CYAN}")
    d.text((1660, 115), "— 策略仓位上限", font=img_font(22), fill=f"#{BLUE}")
    p2 = (left, 690, right, 1030)
    draw_axes(d, p2, 0, 110, [0, 25, 50, 75, 100], months, x_pos)
    bar_w = 78
    for x, value in zip(x_pos, turnover):
        y = p2[3] - value / 110 * (p2[3] - p2[1])
        d.rectangle((x - bar_w / 2, y, x + bar_w / 2, p2[3]), fill=f"#{PURPLE}")
    max_cost = 0.22
    cost_pts = [(x_pos[i], p2[3] - cost[i] / max_cost * (p2[3] - p2[1])) for i in range(len(cost))]
    d.line(cost_pts, fill=f"#{RED}", width=5)
    for x, y in cost_pts:
        d.ellipse((x - 6, y - 6, x + 6, y + 6), fill="white", outline=f"#{RED}", width=3)
    d.text((1320, 625), "■ 单边换手率", font=img_font(22), fill=f"#{PURPLE}")
    d.text((1615, 625), "— 交易成本（右轴缩放）", font=img_font(22), fill=f"#{RED}")
    save_canvas(img, paths["fig10_audit.png"])

    periods_r = ["2022H2", "2023", "2024H1", "2024H2-2025H1"]
    portfolio = [-13.52, -12.57, -9.64, 22.67]
    benchmark_r = [-13.68, -11.38, 0.89, 13.70]
    excess_r = [0.15, -1.19, -10.53, 8.96]
    w, h = 2200, 900
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 38), "跨区间验证：策略没有被包装成“始终跑赢”", font=img_font(44, True), fill=f"#{NAVY}")
    panels = [(120, 200, 1040, 760), (1180, 200, 2100, 760)]
    for p, title in zip(panels, ("不同市场区间累计收益", "相对HS300超额收益")):
        d.text((p[0], 145), title, font=img_font(28, True), fill=f"#{NAVY}")
    p = panels[0]
    zero = p[3] - (0 - (-17)) / (27 - (-17)) * (p[3] - p[1])
    d.line((p[0], zero, p[2], zero), fill="#9BAAB5", width=3)
    centers = [p[0] + 120 + i * 215 for i in range(4)]
    for cx, label, pv, bv in zip(centers, periods_r, portfolio, benchmark_r):
        for offset, val, color in ((-42, pv, CYAN), (42, bv, PURPLE)):
            y = p[3] - (val - (-17)) / 44 * (p[3] - p[1])
            d.rectangle((cx + offset - 34, min(y, zero), cx + offset + 34, max(y, zero)), fill=f"#{color}")
        d.text((cx - text_width(d, label, img_font(18)) / 2, p[3] + 18), label, font=img_font(18), fill=f"#{MUTED}")
    d.text((p[0] + 520, 155), "■ 优化组合", font=img_font(19), fill=f"#{CYAN}")
    d.text((p[0] + 700, 155), "■ 沪深300", font=img_font(19), fill=f"#{PURPLE}")
    p = panels[1]
    zero = p[3] - (0 - (-12)) / (11 - (-12)) * (p[3] - p[1])
    d.line((p[0], zero, p[2], zero), fill="#9BAAB5", width=3)
    centers = [p[0] + 120 + i * 215 for i in range(4)]
    for cx, label, val in zip(centers, periods_r, excess_r):
        y = p[3] - (val - (-12)) / 23 * (p[3] - p[1])
        color = GREEN if val >= 0 else RED
        d.rectangle((cx - 48, min(y, zero), cx + 48, max(y, zero)), fill=f"#{color}")
        d.text((cx - text_width(d, label, img_font(18)) / 2, p[3] + 18), label, font=img_font(18), fill=f"#{MUTED}")
        value_label = f"{val:+.2f}%"
        ly = y - 34 if val >= 0 else y + 8
        d.text((cx - text_width(d, value_label, img_font(19, True)) / 2, ly), value_label, font=img_font(19, True), fill=f"#{color}")
    save_canvas(img, paths["fig11_robustness.png"])

    draw_flow(paths["fig12_validation.png"], "验证证据链", [
        ("单元与接口", "97项Python测试\n全部通过"),
        ("前端静态", "JavaScript语法\n差异格式检查"),
        ("金融不变量", "收益恒等式\nShapley可加性"),
        ("完整运行", "两次32联盟\n真实回测"),
        ("跨区间", "四个市场窗口\n正负结果均保留"),
        ("浏览器验收", "三条曲线\n日期与指标一致"),
    ], [BLUE, CYAN, GREEN, GOLD, PURPLE, RED])
    return paths


def build_document() -> Path:
    demo = json.loads(DEMO_JSON.read_text(encoding="utf-8"))
    figs = make_figures(demo)
    doc = Document()
    style_document(doc)
    add_header_footer(doc)
    bullet_id, decimal_id = add_numbering(doc)

    # Cover
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("ALPHASPLIT"), size=11, bold=True, color=CYAN)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("投研协作激励与贡献计量平台"), size=28, bold=True, color=NAVY)
    p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("完整技术文档 · 初赛提交版"), size=16, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("合理分配激励协作 · 证据可解释 · 决策由评审委员会确认"), size=11, color=MUTED)
    p.paragraph_format.space_after = Pt(34)
    add_figure(doc, figs["fig01_solution.png"], "项目技术闭环概览", width=6.25,
               alt_text="客户描述需求，经Agent确认方案、真实回测、Shapley贡献计量和人工治理应用的五步流程")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    font_run(p.add_run("第五届中国研究生金融科技创新大赛"), size=10.5, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("技术口径日期：2026年8月27日"), size=9.5, color=MUTED)

    # Abstract and contents
    doc.add_page_break()
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_before = Pt(0)
    font_run(h.add_run("摘要"), size=18, bold=True, color=NAVY)
    add_body(doc, "AlphaSplit 是面向资产管理投研协作场景的协作激励、贡献计量与治理支持平台。项目以更合理的分配机制为手段，用可解释贡献证据调动行业、个股、策略、量化、风控五类角色的参与积极性，促进投研协同与决策质量提升，并以长期投资回报改善为目标。历史回测用于形成证据，不构成收益承诺。")
    add_body(doc, "系统把五类角色作为合作博弈参与者，用40个登记因子和少量模块参数表达其决策输入；再通过统一数据、统一成本和统一调仓口径下的32个角色联盟反事实回测，计算各角色对组合相对沪深300超额收益的Shapley边际贡献。")
    add_body(doc, "平台同时引入受约束Agent交互：用户可通过快捷选择和自然语言描述投资偏好，DeepSeek仅负责生成结构化因子配置建议；角色白名单、因子边界、归一化、组合构建、收益计算与贡献计量均由确定性程序完成。配置须经过人工确认才进入实时回测，结果须经过证据白名单校验才进入解读和治理环节。")
    add_body(doc, "本版以当前提交代码中的最终引擎实现为事实基线。官方验证区间为2024年7月至2025年6月：优化策略累计收益22.67%，沪深300累计收益13.70%，超额收益8.96%，夏普比率1.24，最大回撤-5.89%；32个联盟全部完成，Shapley合计与组合超额收益一致。跨区间实验保留了负收益和负超额结果，明确不将单一区间表现解释为未来收益保证。")
    add_callout(doc, "关键词", "投研贡献计量；合作博弈；Shapley值；反事实回测；多智能体规划；因子配置；可解释金融；人工确认；广义价值分配。", fill=LIGHT_GOLD, accent=GOLD)

    doc.add_paragraph(style="Heading 2").add_run("版本事实基线")
    add_table(doc, ["项目", "当前口径", "核验状态"], [
        ["回测引擎", "当前提交代码中的最终实现", "收益、风险与归因口径一致"],
        ["默认配置", "diversified_defensive_v4 / version 4", "五角色各8因子"],
        ["主验证区间", "2024-07-01 至 2025-06-30", "真实月度收益路径"],
        ["对照对象", "优化策略 / 严格旧参数基线 / 沪深300", "同日期、同成本、同基准"],
        ["自动化验证", "97项Python测试 + 前端与浏览器验收", "全部通过"],
    ], [1800, 4400, 3160])

    doc.add_paragraph(style="Heading 2").add_run("目录")
    toc_rows = [
        ["第1章", "问题定义、设计目标与解题思路"],
        ["第2章", "系统架构、角色建模与40因子库"],
        ["第3章", "特征函数、反事实回测与Shapley计量"],
        ["第4章", "Agent规划、客户方案生成与人机协同"],
        ["第5章", "贡献证据、双轨治理与评审委员会决策"],
        ["第6章", "实验结果、收益修复与工程验证"],
        ["附录", "接口、配置、复现命令与参考资料"],
    ]
    add_table(doc, ["章节", "内容"], toc_rows, [1600, 7760])

    # Chapter 1
    chapter_start(doc, "Chapter 01", "第1章  问题定义、设计目标与解题思路",
                  "把传统资产与因子归因推进到组织角色层，以合理分配激励投研协作；在不让大模型生成金融事实的前提下，形成可计量、可解释、可复现的贡献证据。")
    doc.add_paragraph(style="Heading 2").add_run("1.1 应用场景与核心问题")
    add_body(doc, "资产管理产品的结果往往由行业配置、个股研究、仓位择时、量价执行和风险约束共同产生。传统归因可以说明收益来自行业、风格或因子，却难以回答这些决策输入分别由哪些组织角色贡献、角色之间如何协同，以及结果如何进入资源配置和治理讨论。")
    add_body(doc, "AlphaSplit聚焦的不是“给一个最终分配数字”，而是通过可审计的贡献证据形成更合理的分配机制，以调动各角色参与积极性、促进投研协同与决策质量提升，并以长期投资回报改善为目标。分配是组织激励的手段，不是项目的终点。")
    add_body(doc, "证据链回答：角色决策输入如何被因子代理、参数如何进入组合、某角色缺席时组合如何变化、边际贡献如何汇总，以及评审委员会（或机构治理责任人）为何采纳或调整建议。")
    for text in [
        "归因对象缺口：Brinson、Barra和风格归因的对象主要是资产或风险暴露，不直接对应组织角色。",
        "多角色交互：策略、量化和风控常依赖上游选股结果，简单地“去掉一个角色”无法完整分摊交互效应。",
        "自然语言与金融引擎之间存在断层：用户能表达“更稳健”或“重视行业轮动”，但回测需要严格、有限、可校验的参数。",
        "演示与真实性冲突：完整32联盟回测耗时较长，路演需要快速展示，但不能用Mock金融结果冒充真实回测。",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_paragraph(style="Heading 2").add_run("1.2 四层解题框架")
    add_numbered(doc, "资产与市场层：以HS300真实行情、行业映射和公开量价因子构造可复现的投资组合。", decimal_id)
    add_numbered(doc, "决策链映射层：建立“角色—决策职责—因子/参数”映射，使角色参与与否具有可执行含义。", decimal_id)
    add_numbered(doc, "角色贡献层：枚举32个联盟，在控制变量条件下计算特征函数v(S)，再用精确Shapley值分摊交互贡献。", decimal_id)
    add_numbered(doc, "治理应用层：把带符号贡献转为证据轨建议权重，与制度轨按6:4合成；五角色陈述仅供参考，最终由评审委员会独立判断、调整并确认。", decimal_id)
    add_figure(doc, figs["fig01_solution.png"], "图1-1  AlphaSplit解题思路与应用闭环", alt_text="AlphaSplit五步技术闭环")

    doc.add_paragraph(style="Heading 2").add_run("1.3 设计目标")
    add_table(doc, ["目标", "含义", "工程验收"], [
        ["可计量", "输出带符号角色贡献，并能还原组合超额收益", "Σφᵢ = R组合 - RHS300"],
        ["可解释", "每个指标均展示定义、公式、区间、字段和局限", "前端数据依据面板与证据编号"],
        ["可复现", "数据、因子、配置、版本和成本口径可追踪", "request_id、config_id、engine_version"],
        ["可治理", "算法结果是量化建议，不是自动裁决", "角色陈述参考、评审委员会独立确认"],
        ["可演示", "路演可秒开真实预跑结果，且标明来源", "预制包与实时回测明确区分"],
    ], [1500, 4200, 3660])

    doc.add_paragraph(style="Heading 2").add_run("1.4 创新点与边界")
    add_body(doc, "项目的创新不在于单独发明新因子或新归因公式，而在于把角色建模、受约束Agent、反事实回测、Shapley贡献和人工治理连接为可运行系统。归因对象从资产/因子扩展到组织角色；自然语言通过白名单配置进入确定性金融引擎；完整结果通过证据白名单进入解释层。")
    add_callout(doc, "适用边界", "当前直接适用于股票多头、月度调仓和HS300基准场景。40因子是角色决策输入的代理，不等同于真实人员全部能力；单一区间结果不能代表未来表现；系统不替代投资决策、风险审批或治理责任人的判断。", fill=LIGHT_GOLD, accent=GOLD)

    # Chapter 2
    chapter_start(doc, "Chapter 02", "第2章  系统架构、角色建模与40因子库",
                  "系统按交互治理、Agent编排、金融计算、数据审计四层解耦；五角色各配置8个因子，所有配置在进入回测前经过登记、方向和边界校验。")
    doc.add_paragraph(style="Heading 2").add_run("2.1 总体架构")
    add_figure(doc, figs["fig02_architecture.png"], "图2-1  AlphaSplit分层架构", alt_text="AlphaSplit四层软件架构")
    add_body(doc, "交互与治理层由FastAPI静态Web页面承载，提供四步自助流程、三方收益对照、指标依据抽屉、角色贡献陈述参考和评审委员会分配决策。Agent编排层负责意图理解、结构化配置建议、查询规划和证据约束解读。金融计算层完成因子评分、组合构建、成本计算和全联盟Shapley；数据与审计层提供HS300行情、版本标识和逐月诊断。")
    add_table(doc, ["层级", "主要模块", "职责"], [
        ["交互与治理", "engine/web、engine/agentic", "引导流程、图表、角色陈述参考、6:4建议与评审决策"],
        ["Agent编排", "agent_layer、engine/integrations/agent_bridge.py", "规划、因子建议、证据白名单与结果解读"],
        ["金融计算", "backtest/query_engine.py 等", "40因子、五角色Pipeline、交易成本、32联盟Shapley"],
        ["API与保护", "engine/api", "结构化模型、SSE、单任务互斥、可选Token保护"],
        ["数据与场景", "data/processed、data/scenarios", "HS300日线、行业映射、财务与演示场景"],
    ], [1700, 3600, 4060])

    doc.add_paragraph(style="Heading 2").add_run("2.2 五角色职责与组合构建位置")
    add_figure(doc, figs["fig03_role_map.png"], "图2-2  五角色、决策职责与因子类型", alt_text="五角色与40因子映射")
    add_table(doc, ["角色", "主要决策", "在Pipeline中的作用", "模块参数"], [
        ["行业分析师", "识别相对强势行业", "按行业内Top-3股票平均Borda选4个行业", "top_industries=4"],
        ["个股分析师", "形成候选股优先级", "在目标行业中保留30只候选并保留排序", "无"],
        ["策略分析师", "择时与总仓位", "20日趋势为负时启用防御仓位", "target_vol=15%；lookback=20；cap=50%"],
        ["量化分析师", "多策略信号优化", "与个股排序按50%/50%联合Borda选10只", "无"],
        ["风控分析师", "集中度与损失约束", "单股/行业联合容量投影，不把降仓重新拉满", "单股10%；行业30%；止损-5%"],
    ], [1500, 2300, 3400, 2160], small=True)

    doc.add_paragraph(style="Heading 2").add_run("2.3 40因子登记表")
    add_body(doc, "每个角色内部8个因子默认等权12.5%，总和必须归一化为100%；单因子权重上限为35%。direction=+1表示高值优先，direction=-1表示低值优先。反转、突破等基础信号若已在底层完成方向变换，则登记的是最终评分方向。")
    factor_rows = [
        ["行业", "5日动量；10日动量；均线偏离；均线斜率；趋势强度；趋势效率；突破信号；区间位置", "均为高值优先"],
        ["个股", "5日反转；10日反转；RSI强度；RSI反转；随机指标；回撤；下影线支撑；下轨突破", "底层已处理反转方向；最终高值优先"],
        ["策略", "跳过期间动量；跳空溢价；日内收益；已实现波动率；ATR比率；区间波动比；收益偏度；收益峰度", "波动/ATR/区间波动/峰度低值优先"],
        ["量化", "成交量比率；成交量Z分数；成交额活跃度；量价相关性；OBV斜率；收盘时序排名；成交量时序排名；下行波动率", "下行波动率低值优先"],
        ["风控", "5日已实现波动率；10日已实现波动率；下行波动率；ATR比率；区间波动比；回撤；收益偏度；收益峰度", "波动与峰度低值优先；回撤接近0更优"],
    ]
    add_table(doc, ["角色", "8个因子", "方向口径"], factor_rows, [1200, 6100, 2060], small=True)

    doc.add_paragraph(style="Heading 2").add_run("2.4 数据、时间和无前视约束")
    for text in [
        "行情使用前复权日线，组合和HS300基准按照实际共同持有月份对齐。",
        "请求区间开始前额外加载120个自然日，只用于因子与择时预热；收益仍从请求开始日计算。",
        "月末信号用于次月持有，决策时点只读取当时可得数据，不使用未来收益。",
        "财务和行业映射按可得日期处理；任何基准月份缺失都应明确失败，不能静默补零。",
        "缓存键包含engine_version，代码升级后不会复用旧算法缓存。",
    ]:
        add_bullet(doc, text, bullet_id)

    # Chapter 3
    chapter_start(doc, "Chapter 03", "第3章  特征函数、反事实回测与Shapley计量",
                  "把“某角色是否参与”定义为唯一实验变量，在统一数据、组合、成本和基准口径下执行全部角色联盟回测，以精确Shapley分摊交互贡献。")
    doc.add_paragraph(style="Heading 2").add_run("3.1 特征函数定义")
    add_body(doc, "设五角色全集N={行业、个股、策略、量化、风控}。对任意角色子集S，先运行只启用S中角色的完整Pipeline，得到同期组合累计收益R_pipeline(S)；再减去与组合实际持有月份严格对齐的沪深300累计收益R_HS300，定义：")
    add_equation(doc, "v(S) = R_pipeline(S) - R_HS300")
    add_body(doc, "空联盟解释为只持有基准，因此v(∅)=0。全集v(N)等于优化组合相对沪深300的完整超额收益。这样，Shapley有效性对应一个可直接核验的金融恒等式。")
    add_callout(doc, "关键口径", "本版特征函数是“相对HS300的超额收益”，不是组合绝对收益。旧文档中把total_return直接作为v(S)的描述已经失效。", fill=LIGHT_GOLD, accent=GOLD)

    doc.add_paragraph(style="Heading 2").add_run("3.2 角色缺席与控制变量")
    add_body(doc, "角色缺席并不意味着删除一列结果，而是跳过或中性替代其在Pipeline中的决策操作。例如，行业角色缺席时不施加行业选择；个股与量化均缺席时采用deterministic_hash_v1中性选择，避免结果依赖CSV或字典顺序；策略缺席时不主动降仓；风控缺席时不施加集中度投影。")
    add_body(doc, "所有联盟共享同一数据范围、基准、调仓频率、交易成本、候选股票集合和引擎版本，仅参与角色集合不同。任一联盟失败即返回COALITION_BACKTEST_FAILED并终止正式归因，禁止用0填补。")

    doc.add_paragraph(style="Heading 2").add_run("3.3 精确Shapley计算")
    add_equation(doc, "φᵢ = Σ_{S⊆N\\{i}} [ |S|!(n-|S|-1)! / n! ] · [ v(S∪{i}) - v(S) ]")
    add_body(doc, "五角色仅有2⁵=32个联盟，因此本项目采用全枚举而非蒙特卡洛近似。每个角色在所有可能加入顺序中的边际超额贡献被公平平均，角色间交互由Shapley权重自然分摊。")
    add_figure(doc, figs["fig05_shapley.png"], "图3-1  32联盟反事实回测与精确Shapley流程", alt_text="Shapley全联盟回测流程")
    add_table(doc, ["方法", "计算对象", "优点", "局限"], [
        ["完整组合", "v(N)", "直接回答策略整体表现", "不能拆分角色贡献"],
        ["Leave-One-Out", "v(N)-v(N\\{i})", "直观、计算量低", "只看完整组合附近，忽略其他联盟交互"],
        ["精确Shapley", "所有联盟的加权边际贡献", "满足有效性并分摊交互", "计算成本随角色数指数增长"],
    ], [1800, 2300, 3000, 2260])

    doc.add_paragraph(style="Heading 2").add_run("3.4 优化后的组合构建")
    add_figure(doc, figs["fig04_pipeline.png"], "图3-2  月度组合构建、持有与成本审计", alt_text="组合构建流水线")
    add_body(doc, "行业角色先输出全市场股票排序，再按行业内排名最高的3只股票计算平均Borda得分，降低行业成分数量带来的规模偏差。个股角色和量化角色分别转换为排名百分位，以50%/50%联合Borda选出最终10只股票。候选不足时在目标行业内使用确定性中性顺序补足。")
    add_body(doc, "风控不再采用“裁剪后整体归一化”的错误做法，而是进行单股与行业联合容量投影：行业不超过30%，单股不超过10%，相对权重尽可能保留；只有容量确实不足或策略主动防御时才保留现金。交易成本按实际单边换手率计算：")
    add_equation(doc, "Cost_t = (commission_rate + slippage_rate) × one_way_turnover_t")
    add_body(doc, "换手同时考虑股票权重和现金权重变化。初次满仓建仓或完整换仓的单边换手率为100%，持仓不变时为0，不再每月固定扣除成本。")

    doc.add_paragraph(style="Heading 2").add_run("3.5 金融不变量与审计输出")
    for text in [
        "收益恒等式：excess_return = total_return - benchmark_return。",
        "贡献有效性：Σφᵢ = excess_return，允许误差必须在数值容差内。",
        "风险约束：任一股票≤10%，任一行业≤30%；策略降仓不得被风控重新归一化到满仓。",
        "逐月诊断：输出行业、股票数、总仓位、择时仓位、单边换手率和交易成本。",
        "可追溯标识：输出request_id、config_id、engine_version和结果来源。",
    ]:
        add_bullet(doc, text, bullet_id)

    # Chapter 4
    chapter_start(doc, "Chapter 04", "第4章  Agent规划、客户方案生成与人机协同",
                  "Agent承担语义理解、配置建议和证据约束解释；确定性校验器与金融引擎掌握最终事实。所有配置先预览、后确认，失败时明确降级。")
    doc.add_paragraph(style="Heading 2").add_run("4.1 Agent能力定位")
    add_body(doc, "五类投研角色首先是贡献计量对象，不等同于五个自由对话的大模型实例。当前主链中的Agent能力包括：把客户意图转为受约束因子配置建议、把自然语言查询整理为结构化回测计划、把可信证据转为可读解释。大模型不直接生成收益率、风险指标、Shapley值或最终应用比例。")
    add_table(doc, ["Agent能力", "输入", "输出", "禁止事项"], [
        ["因子配置Agent", "偏好选项、自由文字、当前配置", "白名单因子/模块参数补丁", "自由创造因子；绕过边界；自动生效"],
        ["规划Agent", "自然语言或路演预设", "结构化时间、基准、角色与调仓计划", "直接计算金融指标"],
        ["解读Agent", "证据白名单与证据编号", "带依据的摘要、风险提示", "引用未知证据；承诺未来收益"],
        ["角色贡献陈述", "贡献、证据轨、6:4合成建议", "仅供评审参考的角色意见", "参与建议份额计算或最终裁决"],
    ], [1800, 2600, 2800, 2160], small=True)

    doc.add_paragraph(style="Heading 2").add_run("4.2 DeepSeek受约束配置闭环")
    add_figure(doc, figs["fig06_agent.png"], "图4-1  从客户偏好到已确认因子配置", alt_text="DeepSeek因子配置建议与人工确认闭环")
    add_numbered(doc, "用户通过稳健/均衡/进取、控制回撤/获取超额/行业轮动等快捷选项，或自由文字描述需求。", decimal_id)
    add_numbered(doc, "DeepSeek返回结构化建议，只能引用登记角色、40个因子及少量模块参数。", decimal_id)
    add_numbered(doc, "后端执行未知字段拒绝、数值有限性、参数范围、角色内归一化和单因子35%上限校验。", decimal_id)
    add_numbered(doc, "前端展示当前值、建议值、调整理由、模型来源、审计号和配置指纹；配置仍未进入回测。", decimal_id)
    add_numbered(doc, "用户确认后生成config_id和provenance；未修改项保持不变，恢复默认可随时撤销建议。", decimal_id)
    add_numbered(doc, "已确认factor_config随实时回测请求进入确定性引擎；预制模式只回放匹配的真实预跑结果。", decimal_id)

    doc.add_paragraph(style="Heading 2").add_run("4.3 降级、重试与审计")
    add_body(doc, "API Key可在页面中选填，也可通过服务端环境变量DEEPSEEK_API_KEY配置。页面输入的Key仅保存在当前本地服务进程内存中，不写入项目文件、浏览器存储、日志或接口响应，服务重启后自动清除；不输入时默认使用随包提供且重新经过约束校验的离线Agent缓存。模型调用采用有限重试；若JSON Schema响应格式不被支持，则降级到JSON Object，但不会从自由文本中用正则猜测JSON。超时、JSON错误、越界或模型不可用时，系统明确返回rule_fallback及原因，不把规则结果伪装成大模型结果。")
    add_body(doc, "每次建议记录独立audit_id、配置前后指纹、变更数量、来源、耗时和人工确认状态。日志不记录API Key、完整提示词、完整服务地址或原始错误正文，降低敏感信息泄露风险。")
    add_callout(doc, "收敛条件", "Agent流程不是不断改参数直到历史收益最高，而是在配置合法、解释充分且经人工确认后进入回测。反复根据同一区间收益调参会造成样本内过拟合，系统应明确阻止这种叙述。", fill=LIGHT_GOLD, accent=GOLD)

    doc.add_paragraph(style="Heading 2").add_run("4.4 四步自助体验与路演模式")
    add_table(doc, ["步骤", "用户动作", "系统产出", "下一步解锁条件"], [
        ["1 描述需求", "选择偏好或与Agent对话", "结构化配置建议", "建议生成成功"],
        ["2 确认方案", "检查变更与理由并确认", "已确认factor_config", "人工确认完成"],
        ["3 回测理解", "预制快速体验或实时回测", "三方曲线、指标和贡献证据", "回测结果有效"],
        ["4 分配决策", "查看双轨建议与角色参考意见", "评审建议与决策记录", "评审委员会确认"],
    ], [1400, 2800, 3000, 2160])
    add_body(doc, "路演默认推荐“使用预制数据快速体验”。预制包来自同一真实引擎的完整预跑，包含优化策略、严格旧参数基线和HS300同日期月度曲线；它不是Mock，也不会声称回放的是任意现场临时生成配置。实时回测适用于验证已确认的客户专属配置，耗时取决于数据和32联盟计算。")

    doc.add_paragraph(style="Heading 2").add_run("4.5 证据白名单与可解释指标")
    add_body(doc, "回测结果进入解读Agent前，系统检查所有数值有限性、最大回撤范围、收益恒等式、角色权重合计、单角色50%上限和Shapley可加性。不一致字段被隔离并生成警告；可信字段生成E001等证据编号。涉及数字的模型结论必须引用合法证据编号，否则自动降级为确定性解释。")
    add_body(doc, "前端每个核心指标提供“数据依据”：显示定义、公式、本次代入、时间区间、基准、后端字段、request_id、config_id和使用局限。这使页面上的22.67%、8.96%、1.24和-5.89%都能被追溯，而不是孤立展示。")

    # Chapter 5
    chapter_start(doc, "Chapter 05", "第5章  贡献证据、双轨治理与评审委员会决策",
                  "带符号Shapley用于解释角色对超额收益的边际影响；只有正向贡献进入证据轨建议权重，再与制度轨按6:4合成。角色陈述仅供参考，最终由评审委员会独立决策。")
    doc.add_paragraph(style="Heading 2").add_run("5.1 从带符号贡献到证据轨权重")
    add_body(doc, "role_breakdown.shapley_raw保留正负符号。正值说明该角色在所有联盟中的平均边际作用提高了组合超额收益，负值说明在当前区间与当前设定下形成拖累。负贡献不是程序错误，也不应通过更换样本或反复调参强行改成正值。")
    add_body(doc, "证据轨weight仅使用正向Shapley，并设置单角色50%上限。若某角色先达到上限，剩余权重通过水位法分配给其他正贡献角色。因此，Shapley贡献占比与证据轨建议权重不是同一概念。")
    add_figure(doc, figs["fig09_roles.png"], "图5-1  角色带符号贡献与证据轨建议权重", alt_text="五角色Shapley贡献和证据轨权重柱状图")

    doc.add_paragraph(style="Heading 2").add_run("5.2 6:4双轨机制")
    add_equation(doc, "Final_i = 60% × Institutional_i + 40% × Evidence_i")
    add_figure(doc, figs["fig07_governance.png"], "图5-2  制度轨、证据轨与最终合成", alt_text="6比4双轨治理流程")
    add_body(doc, "制度轨内部比例、证据轨内部比例和6:4合成建议必须在前端明确区分。角色贡献陈述以合成建议和贡献证据为参考背景，但不参与建议份额计算；评审委员会（或机构治理责任人）据此独立判断、调整并确认，而不是由角色Agent决定分配。")
    add_table(doc, ["角色", "Shapley贡献", "证据轨内部权重", "制度轨等权时最终比例"], [
        ["行业分析师", "+19.23%", "50.00%", "32.00%"],
        ["个股分析师", "-5.26%", "0.00%", "12.00%"],
        ["策略分析师", "+2.98%", "50.00%", "32.00%"],
        ["量化分析师", "-6.87%", "0.00%", "12.00%"],
        ["风控分析师", "-1.12%", "0.00%", "12.00%"],
        ["合计", "+8.96%", "100.00%", "100.00%"],
    ], [1800, 1900, 2400, 3260])
    add_callout(doc, "解释示例", "行业角色的+19.23%大于组合整体+8.96%，是因为个股、量化和风控的负边际贡献抵消了部分行业与策略的正贡献。这正是保留带符号Shapley而不是只看百分比排名的价值。", fill=LIGHT_BLUE, accent=BLUE)

    doc.add_paragraph(style="Heading 2").add_run("5.3 角色贡献陈述与评审委员会决策")
    add_body(doc, "系统读取角色贡献、证据轨权重和6:4合成建议，组织五类角色基于证据陈述意见。行业角色可以说明其主要正贡献来源；负贡献角色可以提出市场阶段、约束成本或后续复盘计划。这些陈述仅构成评审参考条件，不改变6:4合成建议，也不构成分配决定。")
    add_body(doc, "评审委员会（或机构治理责任人）是分配决策者，可综合量化证据、角色陈述和机构治理规则独立判断、调整并确认。决策记录应包含原方案、纯证据方案、合成建议、角色参考意见、确认人、时间和版本。系统不直接执行预算、研究资源、项目机会、数据/算力额度或其他资源的划拨。")

    doc.add_paragraph(style="Heading 2").add_run("5.4 适用的广义价值分配场景")
    for text in [
        "投研预算：根据角色贡献与制度要求调整下一周期研究预算。",
        "研究资源：分配数据订阅、算力、外部专家和调研额度。",
        "项目机会：决定重点课题、产品孵化和跨团队协作机会。",
        "组织复盘：识别当前市场状态下的正向驱动和负向拖累，形成改进计划。",
        "治理审计：保留算法证据、角色参考意见和评审委员会决策记录，支持事后复核。",
    ]:
        add_bullet(doc, text, bullet_id)

    # Chapter 6
    chapter_start(doc, "Chapter 06", "第6章  实验结果、收益修复与工程验证",
                  "本章报告优化策略、严格旧参数基线和HS300的真实同口径结果，同时解释29.51%与1.43%两种历史结果为何都不能作为当前可信基线。")
    doc.add_paragraph(style="Heading 2").add_run("6.1 主实验设置")
    add_table(doc, ["项目", "设置"], [
        ["请求ID", "optimized_diversified_defensive_v4_20260827"],
        ["配置ID", "diversified_defensive_v4 / version 4"],
        ["区间", "2024-07-01 至 2025-06-30"],
        ["基准", "沪深300，同实际持有月份"],
        ["股票与交易日", "298只股票；241个有效交易日"],
        ["调仓与持仓", "月度调仓；12个月均为10只股票"],
        ["成本", "实际单边换手率 × 0.20%综合费率与滑点"],
        ["贡献计量", "31个非空联盟 + 1个空联盟；精确Shapley"],
    ], [2500, 6860])

    doc.add_paragraph(style="Heading 2").add_run("6.2 三方案表现对比")
    add_figure(doc, figs["fig08_returns.png"], "图6-1  优化策略、严格旧参数基线与沪深300累计收益路径",
               alt_text="2024年7月至2025年6月三条累计收益曲线")
    add_table(doc, ["指标", "优化策略", "严格旧参数基线", "沪深300"], [
        ["累计收益", "22.67%", "9.30%", "13.70%"],
        ["相对沪深300超额", "+8.96%", "-4.41%", "0.00%"],
        ["年化夏普比率", "1.24", "0.54", "0.51"],
        ["最大回撤", "-5.89%", "-9.10%", "-6.15%"],
    ], [2500, 2286, 2287, 2287])
    add_body(doc, "优化策略相对严格旧参数基线累计收益提高13.37个百分点，同时最大回撤从-9.10%收窄到-5.89%。三条曲线均由各自真实月收益连续复利得到，不使用期末指标倒推过程数据。")

    doc.add_paragraph(style="Heading 2").add_run("6.3 收益暴跌的根因与修复")
    add_table(doc, ["阶段", "显示结果", "主要问题", "结论"], [
        ["旧高收益版本", "约29.51%", "10只股票先裁剪至5%，随后整体归一化又回到约10%；风控上限实际失效", "不应作为有效风控下可信基线"],
        ["严格修复初版", "约1.43%", "10只×5%使总股票容量最多50%；另有首月缺失、候选稀疏和固定扣费", "收益被多重实现问题压低"],
        ["当前v7", "22.67%", "10%单股上限、4行业、联合排序、120日预热、实际换手成本", "当前可复现主结果"],
    ], [1700, 1300, 4300, 2060], small=True)
    add_body(doc, "本次修复没有恢复错误归一化，也没有把高收益数字写死。核心做法是让约束参数与组合规模匹配，并修复行业规模偏差、个股/量化职责覆盖、首月预热、候选补足和成本口径。")

    doc.add_paragraph(style="Heading 2").add_run("6.4 月度组合审计")
    add_figure(doc, figs["fig10_audit.png"], "图6-2  逐月仓位、换手率与交易成本", alt_text="12个月组合仓位换手和交易成本")
    add_body(doc, "12个月每月均持有10只股票，实际股票仓位为40%至100%，平均单边换手率约65.62%，模型累计交易成本约1.57%。低于100%的月份来自行业容量、集中度和弱趋势防御，而不是候选股票不足。")

    doc.add_paragraph(style="Heading 2").add_run("6.5 跨区间稳健性")
    add_figure(doc, figs["fig11_robustness.png"], "图6-3  四个市场区间的组合、基准与超额收益",
               alt_text="跨区间收益和超额收益柱状图")
    add_table(doc, ["区间", "优化组合", "沪深300", "超额收益"], [
        ["2022-07 至 2022-12", "-13.52%", "-13.68%", "+0.15%"],
        ["2023-01 至 2023-12", "-12.57%", "-11.38%", "-1.19%"],
        ["2024-01 至 2024-06", "-9.64%", "+0.89%", "-10.53%"],
        ["2024-07 至 2025-06", "+22.67%", "+13.70%", "+8.96%"],
    ], [2800, 2186, 2187, 2187])
    add_callout(doc, "结论边界", "四个区间中只有两个取得正超额，2024年上半年明显跑输。路演区间说明当前方法能够在特定市场状态下取得较好表现，但不能证明策略在所有时期稳定跑赢。", fill=LIGHT_GOLD, accent=GOLD)

    doc.add_paragraph(style="Heading 2").add_run("6.6 角色贡献结果")
    add_body(doc, "主实验中，行业分析师贡献+19.23%，策略分析师贡献+2.98%；个股、量化和风控分别为-5.26%、-6.87%和-1.12%。五项合计+8.96%，与组合超额收益一致。角色贡献随市场阶段变化是正常现象，负贡献应进入复盘和敏感性分析，而不是被隐藏。")
    add_figure(doc, figs["fig09_roles.png"], "图6-4  主实验角色Shapley贡献与证据轨权重",
               alt_text="主实验五角色贡献与证据轨权重")

    doc.add_paragraph(style="Heading 2").add_run("6.7 工程验证")
    add_figure(doc, figs["fig12_validation.png"], "图6-5  从单元测试到浏览器验收的验证链路", alt_text="AlphaSplit验证证据链")
    add_table(doc, ["验证层", "结果", "覆盖重点"], [
        ["Python自动化", "97项全部通过", "因子配置、风控投影、换手成本、Agent、API、产品文案"],
        ["前端静态", "node --check通过", "JavaScript语法；资源引用"],
        ["补丁质量", "git diff --check通过", "无补丁格式错误"],
        ["完整金融运行", "两次32联盟成功", "优化策略与严格旧参数基线；可加性通过"],
        ["跨区间", "四个窗口完成", "正负收益和正负超额均保留"],
        ["浏览器验收", "通过", "默认日期、22.67%、9.30%、13.70%及三条曲线"],
    ], [1900, 2100, 5360])

    doc.add_paragraph(style="Heading 2").add_run("6.8 局限与下一步")
    for text in [
        "角色因子仍是岗位决策代理，需要机构脱敏决策日志进行校准。",
        "当前股票多头和HS300基准适用范围有限，应扩展多资产与其他基准。",
        "完整Shapley随角色数指数增长，角色扩展后需报告近似误差与排序稳定性。",
        "客户专属配置需做滚动窗口、样本外和压力场景验证，不能只看单一区间。",
        "实时回测仍是同步长任务，后续应增加任务队列、取消、超时和进度持久化。",
        "前端应继续增加端到端自动化回归，确保按钮链路、状态恢复和移动端布局长期稳定。",
    ]:
        add_bullet(doc, text, bullet_id)

    # Appendices
    chapter_start(doc, "Appendix A", "附录A  关键接口与数据契约",
                  "接口按规划、配置、回测、解释、角色陈述和评审确认分层；所有写入回测的配置必须经过结构化模型与白名单校验。")
    add_table(doc, ["接口", "作用", "关键输出"], [
        ["GET /api/health", "服务、模式和引擎可用性检查", "status、mode、version"],
        ["GET /api/agent/factor-config", "返回五角色、40因子和默认配置", "roles、default_config、agent"],
        ["POST /api/agent/factor-config", "生成受约束配置变更预览", "changes、config、source、audit"],
        ["POST /api/agent/plan", "自然语言/预设转结构化回测计划", "query、summary、warnings"],
        ["POST /api/agent/interpret", "基于证据白名单解释回测", "summary、evidence_refs"],
        ["POST /api/agent/pipeline", "串联规划、回测和解释", "plan、backtest、interpretation"],
        ["POST /api/backtest/run", "运行已确认配置的完整回测", "overview、role_breakdown、diagnostics"],
        ["GET /api/backtest/demo", "返回真实预跑官方演示包", "comparison、results、source"],
        ["POST /api/negotiate", "生成角色贡献陈述与6:4评审建议（接口名兼容保留）", "claims、proposal、status"],
        ["POST /api/confirm", "评审委员会确认最终分配决定", "confirmed、audit_record"],
    ], [3000, 3300, 3060], small=True)

    doc.add_paragraph(style="Heading 2").add_run("A.1 核心回测结果字段")
    add_table(doc, ["字段", "定义", "当前主实验"], [
        ["total_return", "组合月收益连续复利", "0.226659"],
        ["benchmark_return", "同期HS300月收益连续复利", "0.137050"],
        ["excess_return", "组合累计收益 - 基准累计收益", "0.089610"],
        ["sharpe_ratio", "月收益风险调整后的年化指标", "1.235432"],
        ["max_drawdown", "累计净值相对历史高点的最大跌幅", "-0.058943"],
        ["shapley_raw", "角色对超额收益的带符号平均边际贡献", "五角色合计0.089610"],
        ["weight", "仅正贡献、单角色50%上限后的证据轨权重", "行业50%；策略50%"],
    ], [2500, 4300, 2560])

    doc.add_paragraph(style="Heading 2").add_run("A.2 默认模块参数")
    add_table(doc, ["角色", "参数", "默认值", "允许范围/说明"], [
        ["行业", "top_industries", "4", "2至6"],
        ["策略", "target_vol", "0.15", "0.05至0.30"],
        ["策略", "lookback", "20", "10至60个交易日"],
        ["策略", "defensive_position_cap", "0.50", "0.20至1.00"],
        ["风控", "max_single", "0.10", "0.01至0.20"],
        ["风控", "max_industry", "0.30", "0.10至0.60"],
        ["风控", "stop_loss_monthly", "-0.05", "-0.30至-0.01"],
    ], [1800, 2800, 1800, 2960])

    chapter_start(doc, "Appendix B", "附录B  复现、测试与运行说明",
                  "所有命令均在项目根目录执行。正式路演使用真实预跑演示包；实时回测须先重启服务以加载当前Python代码。")
    doc.add_paragraph(style="Heading 2").add_run("B.1 前台启动")
    add_callout(doc, "PowerShell", "cd F:\\wills\\competitions\\AlphaSplit\n.\\venv\\Scripts\\Activate.ps1\npython -m engine.api.main --mode mock --host 127.0.0.1 --port 8000", fill=LIGHT, accent=BLUE)
    doc.add_paragraph(style="Heading 2").add_run("B.2 自动化测试")
    add_callout(doc, "PowerShell", "python -m unittest discover -s tests -v\nnode --check engine/web/app.js\ngit diff --check", fill=LIGHT, accent=GREEN)
    doc.add_paragraph(style="Heading 2").add_run("B.3 策略诊断与演示包")
    add_callout(doc, "PowerShell", "python scripts/evaluate_backtest_strategy.py\npython scripts/gen_backtest_demo.py", fill=LIGHT, accent=GOLD)
    add_body(doc, "演示包engine/web/demo/backtest_demo.json同时保存优化策略、严格旧参数基线和HS300的月度路径，并记录customer_request_id、default_request_id、config_id与same_scope。前端应标注“真实预跑结果回放”，避免与实时临时配置混淆。")

    doc.add_paragraph(style="Heading 2").add_run("B.4 路演推荐顺序")
    for text in [
        "展示四步主线并说明Agent不直接生成金融数值。",
        "保留默认偏好，生成配置建议，展示调整前后和人工确认门槛。",
        "选择预制数据快速体验，展示优化策略、严格基线和HS300三条曲线。",
        "点击指标依据，解释累计收益、超额、夏普、回撤和Shapley来源。",
        "进入应用建议，明确制度轨60%、证据轨40%和最终合成比例。",
        "进入贡献评审：角色陈述仅供参考，6:4建议不受其影响，最后由评审委员会独立确认。",
    ]:
        add_numbered(doc, text, decimal_id)

    chapter_start(doc, "References", "参考资料与版本来源",
                  "理论依据以公开论文和行业资料为主，实现口径以2026年8月27日随包源码、离线演示数据和自动化测试结果为准。")
    refs = [
        "[1] Shapley, L. S. A Value for n-Person Games. Contributions to the Theory of Games II, 1953, pp. 307-317.",
        "[2] Brinson, G. P., Hood, L. R., and Beebower, G. L. Determinants of Portfolio Performance. Financial Analysts Journal, 1986, 42(4): 39-44.",
        "[3] Ghorbani, A., and Zou, J. Data Shapley: Equitable Valuation of Data for Machine Learning. ICML 2019. https://arxiv.org/pdf/1904.02868",
        "[4] Moehle, N., Boyd, S., and Ang, A. Portfolio Performance Attribution via Shapley Value. 2021. https://arxiv.org/pdf/2102.05799",
        "[5] Mitchell, R., Cooper, J., Frank, E., and Holmes, G. Sampling Permutations for Shapley Value Estimation. Journal of Machine Learning Research, 2022, 23(43): 1-46. https://jmlr.org/papers/v23/21-0439.html",
        "[6] Lundberg, S. M., and Lee, S.-I. A Unified Approach to Interpreting Model Predictions. NeurIPS 2017. https://arxiv.org/pdf/1705.07874",
        "[7] Microsoft Research. RD-Agent: Research and Development Agent. https://github.com/microsoft/RD-Agent (accessed 2026-08-27).",
        "[8] 海通证券研究所.《Brinson模型的升级：风格绩效归因模型》. 2023-06. https://www.htsec.com/jfimg/colimg/upload/20230703/29131688345186124.pdf",
        "[9] 实现核验来源：backtest/query_engine.py、risk_control_module.py、factor_config.py、timing_module.py，以及engine/api/main.py。",
        "[10] 离线演示与复现实例：engine/web/demo/backtest_demo.json；自动化验证见tests/目录。",
    ]
    for item in refs:
        add_numbered(doc, item, decimal_id)
    add_callout(doc, "版本说明", "若本文件与更早章节稿在收益率、风控参数、特征函数、Agent提供商、测试数量或接口字段上存在冲突，以当前提交源码和官方演示包为准。", fill=LIGHT_GOLD, accent=GOLD)

    doc.core_properties.title = "AlphaSplit 投研协作激励与贡献计量平台：完整技术文档"
    doc.core_properties.subject = "金融科技竞赛项目完整技术文档"
    doc.core_properties.author = "AlphaSplit 项目组"
    doc.core_properties.keywords = "AlphaSplit, Shapley, 反事实回测, Agent, 因子配置, 广义价值分配"
    doc.core_properties.comments = "基于当前提交源码与2026-08-27真实回测结果整合"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_document())
