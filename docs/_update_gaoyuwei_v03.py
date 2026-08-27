# -*- coding: utf-8 -*-
"""Update the retained Gaoyuwei chapter draft to the verified Stage-3 project state."""
from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\wills\competitions\AlphaSplit")
SOURCE = Path(r"F:\wills\competitions\材料\AlphaSplit_技术文档_高玉伟章节初稿_v0.1.docx")
OUTPUT = ROOT / "docs" / "AlphaSplit_技术文档_高玉伟章节完善版_v0.3.docx"
RESULT = ROOT / "backtest" / "results" / "stage3_20260825" / "response.json"
ASSET_DIR = ROOT / "docs" / ".qa_gaoyuwei_v03" / "assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

ROLES = ["行业分析师", "个股分析师", "策略分析师", "量化分析师", "风控分析师"]
COLORS = ["#5687d8", "#28a89c", "#9b76d1", "#c7a65a", "#df7d73"]
NARROW_TERMS = ("薪资", "发薪", "工资", "薪酬", "奖金", "落账")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc")
    return ImageFont.truetype(str(path), size=size)


def rounded(draw: ImageDraw.ImageDraw, box, fill, outline=None, radius=18, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw: ImageDraw.ImageDraw, x1, y, x2, color="#4f819c"):
    draw.line((x1, y, x2 - 15, y), fill=color, width=5)
    draw.polygon([(x2 - 15, y - 10), (x2, y), (x2 - 15, y + 10)], fill=color)


def save_pipeline_chart() -> Path:
    out = ASSET_DIR / "fig1-1_pipeline_stage3.png"
    img = Image.new("RGB", (1400, 520), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "AlphaSplit 解题主线", font=font(42, True), fill="#19324a")
    cards = [
        ("需求与 Agent 规划", "自然语言转 query"),
        ("五角色 Pipeline", "行业→个股→量化\n策略→风控"),
        ("32 联盟反事实", "非空联盟真实回测"),
        ("超额收益 Shapley", "严格对齐同期 HS300"),
        ("双轨与人工确认", "制度 60% + 证据 40%"),
    ]
    x0, y0, w, h, gap = 40, 145, 235, 230, 45
    for i, (title, body) in enumerate(cards):
        x = x0 + i * (w + gap)
        rounded(d, (x, y0, x + w, y0 + h), "#ffffff", outline="#b9cbd7", radius=18, width=3)
        d.text((x + 18, y0 + 45), title, font=font(21, True), fill="#203c50")
        for j, line in enumerate(body.split("\n")):
            d.text((x + 18, y0 + 122 + 40 * j), line, font=font(16), fill="#607787")
        if i < len(cards) - 1:
            arrow(d, x + w + 8, y0 + h // 2, x + w + gap - 8)
    d.text((60, 435), "带符号贡献被如实保留；系统提供证据与建议，不替代治理责任人的最终判断。", font=font(22), fill="#36566d")
    img.save(out, quality=95)
    return out


def save_architecture_chart() -> Path:
    out = ASSET_DIR / "fig2-1_architecture_stage3.png"
    img = Image.new("RGB", (1400, 860), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "AlphaSplit 四层系统架构", font=font(42, True), fill="#19324a")
    layers = [
        ("治理输出层", "贡献可视化 · 制度轨/证据轨合成 · 五角色协商 · 人工确认", "#e9f2fb"),
        ("智能层", "agent_layer 规划/解读 · 五角色协商 · Mock/LLM 可替换", "#eaf7f2"),
        ("计量层", "五角色 Pipeline · HS300 对齐 · 32 联盟反事实 · 超额收益 Shapley", "#fff4df"),
        ("信号层", "HS300 行情 · 行业映射 · 角色因子 · 确定性中性选择", "#f1edfb"),
    ]
    y0, h, gap = 125, 145, 35
    for i, (title, body, fill) in enumerate(layers):
        y = y0 + i * (h + gap)
        rounded(d, (60, y, 1340, y + h), fill, outline="#9fb3c2", radius=18, width=2)
        d.text((90, y + 45), title, font=font(27, True), fill="#203c50")
        d.text((330, y + 48), body, font=font(21), fill="#4f6677")
    d.text((60, 795), "核心边界：因子是角色决策输入代理；角色是计量和治理讨论的直接对象。", font=font(22), fill="#36566d")
    img.save(out, quality=95)
    return out


def save_modules_chart() -> Path:
    out = ASSET_DIR / "fig2-2_modules_stage3.png"
    img = Image.new("RGB", (1400, 820), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "仓库模块与主调用关系", font=font(42, True), fill="#19324a")
    boxes = [
        (80, 140, 390, 280, "engine/web", "ECharts 页面与交互"),
        (545, 140, 855, 280, "engine/api", "FastAPI / SSE / 访问保护"),
        (1010, 140, 1320, 280, "engine/integrations", "回测与 Agent 桥接"),
        (80, 405, 390, 555, "agent_layer", "规划 / 解读 · 可替换实现"),
        (545, 405, 855, 555, "backtest", "五角色 Pipeline + Shapley"),
        (1010, 405, 1320, 555, "engine/agentic", "协商、确认与附录 Mock"),
    ]
    for i, (x1, y1, x2, y2, title, body) in enumerate(boxes):
        rounded(d, (x1, y1, x2, y2), "#ffffff", outline="#b9cbd7", radius=18, width=3)
        d.text(((x1 + x2) // 2, y1 + 45), title, font=font(23, True), fill=COLORS[i % 5], anchor="ma")
        d.text(((x1 + x2) // 2, y1 + 100), body, font=font(17), fill="#526b7c", anchor="ma")
    arrow(d, 398, 210, 537)
    arrow(d, 863, 210, 1002)
    for x in (235, 700, 1165):
        d.line((x, 280, x, 405), fill="#4f819c", width=5)
        d.polygon([(x - 10, 390), (x, 405), (x + 10, 390)], fill="#4f819c")
    rounded(d, (80, 645, 1320, 755), "#eef3f6", outline="#c5d2db", radius=16, width=2)
    d.text((700, 680), "data/processed + data/scenarios + backtest/results：输入、配置与可审计结果", font=font(22), fill="#3f5b70", anchor="ma")
    img.save(out, quality=95)
    return out


def save_web_flow_chart() -> Path:
    out = ASSET_DIR / "fig2-3_web_flow_stage3.png"
    img = Image.new("RGB", (1400, 650), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "Web 端完整工作流", font=font(42, True), fill="#19324a")
    steps = [
        ("参数与 Agent 规划", "日期、行业、角色\n自然语言或预设"),
        ("证据驾驶舱", "数据加载、角色选股\n净值与联盟进度"),
        ("结论摘要", "组合/基准/超额\n带符号角色贡献"),
        ("技术明细", "证据轨权重\nShapley 与不变量"),
        ("人工确认", "双轨合成、协商\n治理责任人确认"),
    ]
    x0, y0, w, h, gap = 40, 165, 235, 305, 45
    for i, (title, body) in enumerate(steps):
        x = x0 + i * (w + gap)
        rounded(d, (x, y0, x + w, y0 + h), "#ffffff", outline="#b9cbd7", radius=18, width=3)
        d.text((x + 18, y0 + 40), f"{i + 1}", font=font(25, True), fill=COLORS[i])
        d.text((x + 18, y0 + 98), title, font=font(20, True), fill="#203c50")
        for j, line in enumerate(body.split("\n")):
            d.text((x + 18, y0 + 175 + 42 * j), line, font=font(16), fill="#607787")
        if i < len(steps) - 1:
            arrow(d, x + w + 8, y0 + h // 2, x + w + gap - 8)
    d.text((60, 555), "路演演示读取预计算 Stage-3 结果；实时回测通过 SSE 展示阶段与联盟进度。", font=font(22), fill="#36566d")
    img.save(out, quality=95)
    return out


def save_shapley_chart() -> Path:
    out = ASSET_DIR / "fig5-1_shapley_stage3.png"
    img = Image.new("RGB", (1400, 680), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "全联盟 Shapley 超额贡献计算", font=font(42, True), fill="#19324a")
    cards = [
        ("角色集合 N", "行业 / 个股 / 策略\n量化 / 风控"),
        ("联盟 S", "遍历 2^5=32 个联盟\n非空联盟真实回测"),
        ("特征函数", "v(∅)=0\nv(S)=R(S)-R_HS300"),
        ("Shapley phi_i", "平均全部加入顺序\n中的边际超额贡献"),
    ]
    x0, y0, w, h, gap = 65, 155, 280, 305, 55
    for i, (title, body) in enumerate(cards):
        x = x0 + i * (w + gap)
        rounded(d, (x, y0, x + w, y0 + h), "#ffffff", outline="#b9cbd7", radius=18, width=3)
        d.text((x + 23, y0 + 52), title, font=font(25, True), fill=COLORS[i])
        for j, line in enumerate(body.split("\n")):
            d.text((x + 23, y0 + 150 + 47 * j), line, font=font(17), fill="#526b7c")
        if i < len(cards) - 1:
            arrow(d, x + w + 10, y0 + h // 2, x + w + gap - 10)
    rounded(d, (65, 525, 1335, 625), "#eaf7f2", outline="#a8ccbd", radius=16, width=2)
    d.text((700, 560), "有效性：sum(phi_i) = v(N) = 完整组合收益 - 同期 HS300 收益；任一联盟失败即中止。", font=font(21, True), fill="#285b4b", anchor="ma")
    img.save(out, quality=95)
    return out


def save_loo_comparison_chart() -> Path:
    out = ASSET_DIR / "fig5-3_loo_shapley_stage3.png"
    img = Image.new("RGB", (1400, 680), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "LOO 与全联盟 Shapley：用途不同，不可混用", font=font(42, True), fill="#19324a")
    panels = [
        (70, "Leave-One-Out", "固定其他角色，移除角色 i\n观察一次边际变化\n约 O(n) 次评估\n适合快速敏感性分析", "#fff4df", "附录/辅助"),
        (730, "全联盟 Shapley", "覆盖全部加入顺序与联盟\n平均角色边际超额贡献\n精确计算 O(2^n)\n满足有效性等公理", "#eaf7f2", "当前主线"),
    ]
    for x, title, body, fill, badge in panels:
        rounded(d, (x, 145, x + 600, 545), fill, outline="#b9cbd7", radius=18, width=3)
        d.text((x + 35, 195), title, font=font(30, True), fill="#203c50")
        rounded(d, (x + 420, 188, x + 555, 235), "#ffffff", outline="#9fb3c2", radius=20, width=2)
        d.text((x + 487, 211), badge, font=font(17, True), fill="#36566d", anchor="mm")
        for j, line in enumerate(body.split("\n")):
            d.text((x + 40, 285 + 52 * j), line, font=font(20), fill="#526b7c")
    d.text((60, 610), "本文定量结论仅使用全联盟 Shapley；LOO 只作为方法对照和敏感性工具。", font=font(22), fill="#36566d")
    img.save(out, quality=95)
    return out


def save_dual_track_chart(evidence: dict[str, float]) -> Path:
    out = ASSET_DIR / "fig5-4_dual_track_stage3.png"
    img = Image.new("RGB", (1400, 760), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "6:4 双轨合成：制度轨 + 证据轨", font=font(42, True), fill="#19324a")
    d.text((70, 105), "默认制度轨各 20%；证据轨来自正向 Shapley 超额贡献并执行单角色 50% 上限", font=font(24), fill="#597083")

    left, top, chart_w, row_h = 345, 190, 900, 92
    for i, role in enumerate(ROLES):
        y = top + i * row_h
        d.text((70, y + 21), role, font=font(25, True), fill="#253b4c")
        institution = 20.0
        ev = evidence[role] * 100
        final = 0.6 * institution + 0.4 * ev
        for value, color, yoff in [(institution, "#9aa9b6", 8), (ev, COLORS[i], 33), (final, "#173f5f", 58)]:
            width = int(chart_w * value / 55.0)
            rounded(d, (left, y + yoff, left + width, y + yoff + 18), color, radius=9, width=1)
        d.text((1265, y + 8), f"制度 {institution:.1f}%", font=font(18), fill="#6f7e89", anchor="ra")
        d.text((1265, y + 34), f"证据 {ev:.1f}%", font=font(18), fill=COLORS[i], anchor="ra")
        d.text((1265, y + 60), f"合成 {final:.1f}%", font=font(18, True), fill="#173f5f", anchor="ra")

    d.text((70, 675), "最终份额 = 60% × 制度轨现行份额 + 40% × 证据轨建议权重；应用前必须人工确认。", font=font(22), fill="#36566d")
    img.save(out, quality=95)
    return out


def save_negotiation_flow() -> Path:
    out = ASSET_DIR / "fig5-6_negotiation_flow.png"
    img = Image.new("RGB", (1400, 620), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "从回测证据到人工确认", font=font(42, True), fill="#19324a")
    steps = [
        ("1", "Shapley 证据", "32 联盟反事实回测\n输出带符号超额贡献"),
        ("2", "双轨合成", "制度轨 60% + 证据轨 40%\n形成可讨论锚点"),
        ("3", "五角色协商", "行业 / 个股 / 策略 / 量化 / 风控\n在带宽内陈述主张"),
        ("4", "人工确认", "治理责任人复核并确认\n系统不自动应用"),
    ]
    box_w, gap, x0, y0 = 275, 55, 60, 180
    for i, (num, title, body) in enumerate(steps):
        x = x0 + i * (box_w + gap)
        rounded(d, (x, y0, x + box_w, y0 + 300), "#ffffff", outline="#b9cbd7", radius=18, width=3)
        rounded(d, (x + 22, y0 + 22, x + 76, y0 + 76), COLORS[min(i, 4)], radius=27)
        d.text((x + 49, y0 + 49), num, font=font(24, True), fill="white", anchor="mm")
        d.text((x + 25, y0 + 108), title, font=font(29, True), fill="#203c50")
        lines = body.split("\n")
        for j, line in enumerate(lines):
            d.text((x + 25, y0 + 175 + j * 45), line, font=font(20), fill="#607787")
        if i < len(steps) - 1:
            ax = x + box_w + 12
            d.line((ax, y0 + 150, ax + 31, y0 + 150), fill="#4f819c", width=5)
            d.polygon([(ax + 31, y0 + 140), (ax + 47, y0 + 150), (ax + 31, y0 + 160)], fill="#4f819c")
    d.text((70, 540), "原则：证据提供可解释锚点，协商处理制度约束，最终决策由人承担。", font=font(23), fill="#36566d")
    img.save(out, quality=95)
    return out


def save_stage3_result_chart(contrib: dict[str, float], evidence: dict[str, float]) -> Path:
    out = ASSET_DIR / "fig5-7_stage3_results.png"
    img = Image.new("RGB", (1400, 820), "#f7f9fb")
    d = ImageDraw.Draw(img)
    d.text((70, 42), "Stage-3 真实回测：角色超额贡献与证据轨建议权重", font=font(39, True), fill="#19324a")
    d.text((70, 98), "2024H1 · HS300 · 297 只股票 · 117 个交易日 · 32 个联盟 · 联盟失败 0", font=font(22), fill="#597083")

    zero_x, scale = 350, 44
    top, row_h = 185, 102
    d.line((zero_x, top - 20, zero_x, top + row_h * 5 - 25), fill="#92a6b4", width=2)
    for i, role in enumerate(ROLES):
        y = top + i * row_h
        val = contrib[role] * 100
        ev = evidence[role] * 100
        d.text((70, y + 15), role, font=font(25, True), fill="#253b4c")
        if val >= 0:
            x1, x2 = zero_x, zero_x + int(val * scale)
        else:
            x1, x2 = zero_x + int(val * scale), zero_x
        rounded(d, (min(x1, x2), y + 10, max(x1, x2), y + 48), COLORS[i] if val >= 0 else "#d75a5a", radius=10)
        d.text((max(x1, x2) + 12 if val >= 0 else min(x1, x2) - 12, y + 29), f"{val:+.2f}%", font=font(20, True), fill="#284558", anchor="lm" if val >= 0 else "rm")
        evidence_end = 1040 + int(ev * 4.2)
        rounded(d, (1040, y + 11, evidence_end, y + 47), "#173f5f", radius=10)
        d.text((evidence_end + 15, y + 29), f"证据轨 {ev:.1f}%", font=font(20, True), fill="#173f5f", anchor="lm")
    d.text((70, 724), "Shapley 超额贡献合计 +21.85%，等于组合相对 HS300 的完整超额收益；策略角色负值被如实保留。", font=font(22), fill="#36566d")
    img.save(out, quality=95)
    return out


def set_run_font(run, name="宋体", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    if color:
        run.font.color.rgb = color


def replace_paragraph(paragraph, text: str):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def clear_paragraph(paragraph):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)


def replace_image(paragraph, path: Path, width: float):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def style_new_paragraph(paragraph, heading_level: int | None = None, note: bool = False):
    paragraph.paragraph_format.space_after = Pt(6)
    if heading_level is None:
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.first_line_indent = Inches(0.29)
        for run in paragraph.runs:
            set_run_font(run, "宋体", 12, False, RGBColor(0x66, 0x66, 0x66) if note else None)
    else:
        for run in paragraph.runs:
            set_run_font(run, "黑体", 13 if heading_level == 2 else 16, True)


def add_before(target, text: str, style=None, note=False):
    p = target.insert_paragraph_before(text, style=style)
    level = 2 if style == "Heading 2" else None
    style_new_paragraph(p, heading_level=level, note=note)
    return p


def fill_table_cell(cell, text: str, bold=False):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    set_run_font(run, "宋体", 10.5, bold)


def set_table(table, headers, rows):
    while len(table.rows) < len(rows) + 1:
        table.add_row()
    while len(table.rows) > len(rows) + 1:
        table._tbl.remove(table.rows[-1]._tr)
    for i, value in enumerate(headers):
        fill_table_cell(table.rows[0].cells[i], value, True)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    for existing in tr_pr.findall(qn("w:tblHeader")):
        tr_pr.remove(existing)
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            fill_table_cell(table.rows[r].cells[c], str(value))


def add_results_table_before(doc: Document, target, breakdown):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    set_table(
        table,
        ["角色", "Shapley 超额贡献", "证据轨建议权重", "默认 6:4 合成份额"],
        [
            [
                row["role"],
                f'{row["shapley_raw"]:+.4%}',
                f'{row["weight"]:.2%}',
                f'{0.6 * 0.2 + 0.4 * row["weight"]:.2%}',
            ]
            for row in breakdown
        ],
    )
    target._p.addprevious(table._tbl)
    return table


def main():
    raw = json.loads(RESULT.read_text(encoding="utf-8"))
    results = raw["results"]
    overview = results["overview"]
    breakdown = results["role_breakdown"]
    contrib = {row["role"]: float(row["shapley_raw"]) for row in breakdown}
    evidence = {row["role"]: float(row["weight"]) for row in breakdown}

    dual_chart = save_dual_track_chart(evidence)
    negotiation_flow = save_negotiation_flow()
    result_chart = save_stage3_result_chart(contrib, evidence)
    pipeline_chart = save_pipeline_chart()
    architecture_chart = save_architecture_chart()
    modules_chart = save_modules_chart()
    web_flow_chart = save_web_flow_chart()
    shapley_chart = save_shapley_chart()
    loo_chart = save_loo_comparison_chart()

    doc = Document(SOURCE)
    original = list(doc.paragraphs)

    replacements = {
        3: "技术文档 · 分章节完善稿",
        8: "版本：v0.3（对齐 Stage-3 工程、超额收益 Shapley、制度轨/证据轨与 53 项回归测试）",
        13: "本分册覆盖技术文档分工中高玉伟主笔/收口内容。直接计量对象为投研五角色（行业、个股、策略、量化、风控分析师），因子是角色决策输入的可观测代理。主线回测引擎以五角色为联盟成员，遍历 2^5=32 个联盟，以相对同期 HS300 的超额收益为特征函数计算 Shapley 边际贡献。价值应用采用制度轨 60% + 证据轨 40% 的双轨机制，并经五角色 Agent 协商与人工确认形成广义价值分配建议。附录 Mock 场景仅作机制演示，不与主线实证结果混用。",
        14: "本文中的架构图与流程图用于说明机制；第5.6节数值表和 Stage-3 结果图取自当前完整真实回测，作为定量结论依据。",
        16: "【说明】以下定量结果来自 stage3_20260825 完整真实回测；若第6章更换样本区间，应同步更新本章所有结果与图表。",
        18: "AlphaSplit 构建“多智能体驱动的投研角色贡献计量平台”：以五类投研角色为计量对象，在沪深300成分股与公开行情上运行多角色 Pipeline（行业配置→个股候选→量化优化→策略择时→风险控制），通过 32 个角色联盟的反事实回测计算 Shapley 超额贡献。系统将带符号贡献作为可解释证据，不承诺每个角色在任意样本期均为正；证据进入 6:4 双轨后，由五角色 Agent 协商并交由治理责任人确认，定位是决策支持而非自动决策。",
        19: "当前工程已形成可本地运行闭环：FastAPI + ECharts Web 页面、HS300 实时/缓存回测、SSE 进度流、可替换 Agent 层规划与解读、双轨合成与五角色协商。Stage-3 回测覆盖 2024-01-01 至 2024-06-30、297 只股票、117 个交易日和 32 个联盟：组合累计收益 29.51%，同期 HS300 收益 7.66%，超额收益及 Shapley 合计均为 21.85%，有效性校验通过，联盟失败数为 0。",
        26: "系统按信号、计量、智能和治理输出四层组织。信号层沉淀行情、行业映射和五角色因子；计量层执行候选池构建、组合回测及全联盟 Shapley 超额贡献归因；智能层负责自然语言规划、结果解读与多角色协商；治理输出层形成制度轨/证据轨双轨方案，并要求人工确认。",
        31: "（2）计量层：backtest/query_engine.py 对每个非空联盟 S 定义 v(S)=R_pipeline(S)-R_HS300(aligned months)，空联盟表示仅持有基准，故 v(∅)=0。",
        33: "（4）治理输出层：Web 展示带符号贡献、制度轨/证据轨 6:4 合成、带宽内建议案与人工确认。",
        34: "仓库关键模块包括 engine/api（FastAPI 与访问保护）、engine/web（ECharts 路演前端）、engine/integrations（回测与 Agent 桥接）、backtest（五角色 Pipeline 与 Shapley 引擎）、agent_layer（可替换规划/解读实现）、engine/agentic（协商与附录 Mock）以及 data/scenarios。依赖由 requirements-lock.txt 锁定，安装与前台启动职责分离。",
        38: "用户/Agent 规划 query → 结构化参数校验 → POST /api/backtest/run 或 /stream → 五角色 Pipeline 回测 → 31 个非空联盟 + 空联盟反事实 → Shapley 超额贡献 → 证据轨建议权重 → 6:4 双轨合成 → 五角色协商 → 人工确认。",
        41: "可运行性与安全性：官方演示包可离线秒开；默认 Mock 不依赖外部 LLM，可选 DeepSeek 兼容接口；同一进程限制单个真实回测任务；公开演示模式可关闭历史结果与真实回测接口；配置 API Token 后，真实回测、Agent Pipeline 和历史结果均要求 X-AlphaSplit-Token。",
        45: "图2-3  Web 首页与 Agent 规划入口（界面结构示意）",
        47: "2.3–2.5 由董斌/徐远主笔。统稿时需保证角色集合与第5章输出维度一致，并明确因子是角色决策输入代理。JSON 中 role_weights、old_pool、new_pool 等历史字段为兼容接口，正文统一称制度调节系数、制度轨和证据轨。",
        50: "【说明】本章公式、特征函数、角色链路和 Stage-3 数值均已按当前 query_engine 实现与回测输出核对。",
        52: "设角色集合 N={行业,个股,策略,量化,风控}，|N|=5。对任意非空子集 S⊆N，特征函数定义为 v(S)=R_pipeline(S)-R_HS300(S 对应实际持有月份)；空联盟解释为仅持有基准，故 v(∅)=0。角色 i 的 Shapley 值 φ_i 是其在全部加入顺序中的平均边际超额贡献，并满足有效性、对称性和哑元性。",
        54: "工程实现 query_engine._compute_shapley_attribution 遍历 32 个联盟，每个非空联盟执行真实 Pipeline 回测；任一联盟失败即中止并返回 COALITION_BACKTEST_FAILED，不再以 0 填补。输出 role_breakdown 中 shapley_raw 为带符号超额贡献，weight 为仅使用正向 Shapley、经 50% 单角色上限和水位法重分配后的证据轨建议权重。",
        58: "五角色全子集 Shapley 包含 31 次非空联盟评估和 1 个空联盟基准。工程通过复用全联盟结果、参数缓存和 SSE 进度流降低重复计算与等待成本。角色数继续增加时，可扩展 Monte Carlo 排列采样或截断近似，但必须同时报告样本量、误差和排序稳定性。",
        60: "Leave-One-Out 只考察移除单个角色的影响，计算成本较低，适合敏感性分析，但不能替代主线全子集 Shapley。POST /api/attribution 提供的附录 Mock 用于演示另一类机制归因，其计量对象、数据来源和结论均不得与 HS300 五角色真实回测混用。",
        64: "系统采用双轨机制：制度轨占 60%，保留机构现行治理方案；证据轨占 40%，根据正向 Shapley 超额贡献形成建议权重。角色最终建议份额为 60%×制度轨现行份额 + 40%×证据轨建议权重。负贡献角色证据轨权重可为 0，但最终合成份额通常不为 0。例如制度轨各角色均为 20% 时，策略角色最终份额仍为 12%。",
        67: "图5-4  Stage-3 制度轨/证据轨 6:4 合成示意",
        69: "人工确认页由 engine/agentic/role_negotiation.py 驱动五角色协商。各角色引用 Shapley 超额贡献、证据轨权重和 6:4 合成锚点陈述主张，系统在允许带宽内形成建议案，治理责任人决定是否采纳。协商对象与回测 role_breakdown 一一对应。",
        70: "回测规划与解读由独立 agent_layer 实现，并经 engine/integrations/agent_bridge.py 作为唯一桥接入口。默认规则型 Mock 可把自然语言转换为结构化 query；设置 AGENT_LAYER_IMPL 可整体替换实现，配置 DeepSeek 兼容环境变量后可启用 LLM 模式。",
        72: "图5-5  可替换 Agent 层：规划回测面板（界面结构示意）",
        74: "图5-6  五角色协商、双轨合成与人工确认流程",
        77: "Stage-3 完整真实回测结果：组合累计收益 29.5137%，HS300 同期收益 7.6604%，超额收益 21.8533%，Shapley 合计 21.8533%，有效性校验通过。四个角色贡献为正；策略角色因 2024 年 2 月事前防御性降仓而少赚，贡献为 -0.5627%，该负值被如实保留。证据轨权重合计 100%，单角色最大权重为 50%。",
        79: "图5-7  Stage-3 真实回测角色超额贡献与证据轨建议权重",
        85: "（1）直接计量对象为五投研角色，因子是角色决策输入代理。",
        86: "（2）主线使用全联盟 Shapley，特征函数统一为相对同期 HS300 的超额收益。",
        87: "（3）双轨采用制度轨 60% + 证据轨 40%，证据渐进进入治理流程。",
        88: "（4）Agent 默认 Mock、LLM 可选；任何建议均须人工确认。",
        89: "（5）不强制所有角色正贡献，不通过更换样本或反复调参美化结果。",
        91: "当前局限包括：角色因子仍是岗位决策代理；联盟反事实依赖中性替代策略；单一 2024H1 区间不能代表所有市场状态；真实回测为同步长任务；点击级浏览器自动化尚未接入持续回归；机构脱敏决策日志仍待获取。",
        92: "后续将开展滚动窗口与样本外验证，比较确定性哈希、全池等权和行业内等权等中性策略，报告贡献均值、正贡献比例与置信区间；同时评估任务队列、取消/超时机制、字段版本迁移和端到端浏览器测试。",
    }
    for idx, text in replacements.items():
        replace_paragraph(original[idx], text)

    replace_image(original[20], pipeline_chart, 6.2)
    replace_image(original[27], architecture_chart, 6.0)
    replace_image(original[35], modules_chart, 6.0)
    replace_image(original[44], web_flow_chart, 6.0)
    replace_image(original[55], shapley_chart, 5.9)
    replace_image(original[61], loo_chart, 5.9)
    replace_image(original[66], dual_chart, 6.0)
    replace_image(original[73], negotiation_flow, 6.0)
    replace_image(original[78], result_chart, 6.0)

    set_table(
        doc.tables[0],
        ["类别", "组件", "用途"],
        [
            ["语言", "Python 3.11（已验证）", "回测、API 与 Agent 主业务"],
            ["API", "FastAPI / Uvicorn", "结构化 /api/*、SSE 与静态页"],
            ["回测", "backtest/query_engine.py", "HS300 五角色 Pipeline + 超额收益 Shapley"],
            ["前端", "HTML/CSS/JS + ECharts", "贡献看板、双轨方案与人工确认"],
            ["Agent", "agent_layer（Mock / 可选 DeepSeek）", "规划 query、解读结果、协商桥接"],
            ["数据", "data/processed/*.csv", "HS300 行情、行业映射与因子"],
            ["工程", "requirements-lock.txt + tests", "依赖锁定、接口保护与回归验证"],
        ],
    )
    set_table(
        doc.tables[1],
        ["轨道", "比例", "含义", "数据来源"],
        [
            ["制度轨", "60%", "机构现行治理方案", "用户录入；演示默认各角色 20%"],
            ["证据轨", "40%", "正向 Shapley 形成的建议权重", "role_breakdown.weight"],
            ["合成", "100%", "广义价值分配讨论基准", "双轨合成并经人工确认"],
        ],
    )
    set_table(
        doc.tables[2],
        ["接口", "作用"],
        [
            ["POST /api/backtest/run", "HS300 五角色回测与超额收益 Shapley"],
            ["POST /api/backtest/run/stream", "SSE 推送阶段、联盟与完成事件"],
            ["GET /api/backtest/demo", "读取 Stage-3 官方演示包"],
            ["POST /api/agent/plan", "自然语言或预设生成结构化回测计划"],
            ["POST /api/agent/interpret", "解读回测证据并生成报告"],
            ["POST /api/negotiate", "五角色基于证据与制度轨进行协商"],
            ["POST /api/confirm", "人工确认五角色建议方案"],
        ],
    )
    set_table(
        doc.tables[3],
        ["阶段", "事项", "产出"],
        [
            ["工程骨架", "FastAPI + Web Demo", "可本地前台启动"],
            ["归因修正", "超额收益 Shapley + 失败中止", "32 联盟有效性校验"],
            ["角色链路", "个股候选池 + 量化独立优化", "4/5 角色正贡献"],
            ["策略优化", "趋势确认 + 波动率控制", "避免上涨期机械降仓"],
            ["双轨治理", "制度轨/证据轨 6:4 合成", "广义价值分配建议"],
            ["访问保护", "公开演示、Token、单任务互斥", "API 工程保护"],
            ["测试回归", "53 项测试 + 编译/JS/格式检查", "自动化验证通过"],
            ["文档 v0.3", "口径、公式、结果和图表更新", "本完善版"],
        ],
    )
    set_table(
        doc.tables[4],
        ["图号", "文件", "用途"],
        [
            ["图1-1", pipeline_chart.name, "解题总流程"],
            ["图2-1", architecture_chart.name, "总体架构"],
            ["图2-2", modules_chart.name, "模块关系"],
            ["图2-3", web_flow_chart.name, "Web 工作流"],
            ["图4-3", "fig4-3_mock.png", "Mock/LLM 模式"],
            ["图5-1", shapley_chart.name, "Shapley 机制"],
            ["图5-3", loo_chart.name, "LOO 方法对照"],
            ["图5-4", dual_chart.name, "双轨合成"],
            ["图5-5", "fig5-4e_agent_panel.png", "Agent 规划"],
            ["图5-6", negotiation_flow.name, "协商与人工确认"],
            ["图5-7", result_chart.name, "Stage-3 回测结果"],
        ],
    )

    add_results_table_before(doc, original[78], breakdown)

    image_alt_texts = [
        "AlphaSplit 从角色决策输入、联盟回测、Shapley 归因到人工确认的总体流程",
        "AlphaSplit 信号层、计量层、智能层和治理输出层的四层架构",
        "项目代码仓库各模块及其调用关系",
        "默认 Mock 与可选大模型两种 Agent 运行模式",
        "HS300 岗位贡献页面与 Agent 规划入口的界面结构",
        "从角色联盟边际贡献计算 Shapley 超额贡献的过程",
        "Leave-One-Out 与全联盟 Shapley 方法对照",
        "Stage-3 制度轨、证据轨和 6:4 最终合成份额对比",
        "可替换 Agent 层规划回测的界面结构",
        "回测证据、双轨合成、五角色协商到人工确认的流程",
        "Stage-3 五角色 Shapley 超额贡献与证据轨建议权重",
    ]
    for inline, alt in zip(doc.inline_shapes, image_alt_texts, strict=True):
        inline._inline.docPr.set("descr", alt)
        inline._inline.docPr.set("title", alt)

    chapter7 = original[81]
    add_before(chapter7, "5.7 正确性控制与工程验证", style="Heading 2")
    add_before(chapter7, "归因正确性由三组不变量共同约束：第一，Shapley 合计必须等于完整组合相对同期 HS300 的超额收益；第二，基准月份必须与组合实际持有月份严格对齐，缺月即返回 BENCHMARK_DATA_MISSING；第三，任一联盟执行失败即终止正式归因，禁止用 0 填补。")
    add_before(chapter7, "反事实公平性方面，两个选股角色均缺席时采用 deterministic_hash_v1 中性选择，结果不依赖 CSV 或字典载入顺序；负贡献角色保留带符号证据，但不进入证据轨正向权重。API 使用结构化模型拒绝未知字段和倒置日期，并提供单任务互斥、公开演示限制与可选 Token 保护。")
    add_before(chapter7, "截至本版，53 项单元与接口测试全部通过，Python 编译检查、前端 JavaScript 语法检查和 Git 差异格式检查通过；官方演示包通过 Shapley 有效性、权重归一、50% 上限和负贡献不奖励等金融不变量测试。前端点击级自动化仍是待补项，不在本文中声明为已通过。", note=True)

    add_before(chapter7, "5.8 角色链路优化与结果解释", style="Heading 2")
    add_before(chapter7, "个股分析师由直接输出最终 10 只持仓改为提供 30 只候选股；量化分析师使用自身量化因子从候选池优化到最终 10 只。量化缺席时保留个股排名，两者均缺席时才使用中性选择。该修复消除了原实现把“个股分析师”角色名误传给量化优化器的问题，使两个角色承担独立、可检验的边际职责。")
    add_before(chapter7, "策略分析师采用趋势确认与目标波动率结合的事前规则：历史趋势非负或样本不足时保持满仓，仅在负趋势且高波动时降低仓位。2024H1 中该角色只在 2 月降仓至 85.98%，但组合随后上涨，因此少赚并产生 -0.56% 的 Shapley 超额贡献。这是合理信号在当前样本内失效，而非程序故障；继续调参直到转正会造成样本内美化。")
    page_break = chapter7.insert_paragraph_before()
    page_break.add_run().add_break(WD_BREAK.PAGE)

    final_text = "\n".join(p.text for p in doc.paragraphs)
    final_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    found = [term for term in NARROW_TERMS if term in final_text]
    if found:
        raise RuntimeError(f"narrow allocation terms remain: {found}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"saved {OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
