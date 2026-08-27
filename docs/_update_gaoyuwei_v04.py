# -*- coding: utf-8 -*-
"""Build the v0.4 technical chapter from the verified v0.3 document."""
from __future__ import annotations

import importlib.util
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(r"F:\wills\competitions\AlphaSplit")
SOURCE = ROOT / "docs" / "AlphaSplit_技术文档_高玉伟章节完善版_v0.3.docx"
OUTPUT = ROOT / "docs" / "AlphaSplit_技术文档_高玉伟章节完善版_v0.4.docx"
ASSET_DIR = ROOT / "docs" / ".qa_gaoyuwei_v04" / "assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

spec = importlib.util.spec_from_file_location("v03_helpers", ROOT / "docs" / "_update_gaoyuwei_v03.py")
helpers = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(helpers)

NARROW_TERMS = ("薪资", "发薪", "工资", "薪酬", "奖金", "落账")
ROLES = ["行业分析师", "个股分析师", "策略分析师", "量化分析师", "风控分析师"]


def paragraph_by_prefix(doc: Document, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise RuntimeError(f"paragraph not found: {prefix}")


def replace_by_prefix(doc: Document, prefix: str, text: str):
    helpers.replace_paragraph(paragraph_by_prefix(doc, prefix), text)


def fit_text(draw: ImageDraw.ImageDraw, text: str, box_width: int, size: int, bold: bool = False):
    """Draw-ready line wrapping for Chinese text and technical tokens."""
    lines: list[str] = []
    current = ""
    for char in text:
        trial = current + char
        if draw.textbbox((0, 0), trial, font=helpers.font(size, bold))[2] <= box_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def save_overview_flow() -> Path:
    out = ASSET_DIR / "fig1-1_customer_agent_pipeline_v04.png"
    img = Image.new("RGB", (1400, 700), "#f6f8fa")
    d = ImageDraw.Draw(img)
    d.text((60, 36), "AlphaSplit：客户专属方案到可解释应用建议", font=helpers.font(40, True), fill="#19324a")
    steps = [
        ("1", "描述投资需求", "快捷选择 + 自由文字\n风险偏好、目标与关注角色", "#e8f1fb"),
        ("2", "确认 Agent 方案", "DeepSeek / 规则降级\n白名单补丁、约束校验、人工确认", "#e8f6f1"),
        ("3", "回测并理解结果", "客户方案 / 默认方案 / HS300\n同区间、同基准、同调仓频率", "#fff2df"),
        ("4", "形成应用建议", "Shapley 证据 + 6:4 双轨\n五角色协商、最终人工确认", "#f1ecf8"),
    ]
    x0, y0, w, h, gap = 55, 160, 285, 340, 50
    for i, (num, title, body, fill) in enumerate(steps):
        x = x0 + i * (w + gap)
        helpers.rounded(d, (x, y0, x + w, y0 + h), fill, outline="#9fb3c2", radius=16, width=3)
        helpers.rounded(d, (x + 22, y0 + 22, x + 82, y0 + 82), helpers.COLORS[i], radius=30)
        d.text((x + 52, y0 + 52), num, font=helpers.font(25, True), fill="white", anchor="mm")
        d.text((x + 22, y0 + 118), title, font=helpers.font(25, True), fill="#203c50")
        for j, line in enumerate(body.split("\n")):
            for k, wrapped in enumerate(fit_text(d, line, w - 44, 18)):
                d.text((x + 22, y0 + 190 + 38 * (j + k)), wrapped, font=helpers.font(18), fill="#526b7c")
        if i < len(steps) - 1:
            helpers.arrow(d, x + w + 9, y0 + h // 2, x + w + gap - 9)
    helpers.rounded(d, (55, 555, 1345, 645), "#ffffff", outline="#c6d2da", radius=14, width=2)
    d.text((700, 600), "边界：Agent 只生成建议；未经人工确认的配置不进入客户专属回测，回测结果不构成收益承诺。", font=helpers.font(21, True), fill="#36566d", anchor="mm")
    img.save(out, quality=95)
    return out


def save_factor_agent_flow() -> Path:
    out = ASSET_DIR / "fig5-8_factor_agent_v04.png"
    img = Image.new("RGB", (1400, 780), "#f6f8fa")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "客户专属 40 因子方案：受约束的 Agent 调整闭环", font=helpers.font(39, True), fill="#19324a")
    roles = [("行业", "趋势/动量"), ("个股", "反转/形态"), ("策略", "收益/波动"), ("量化", "量价/流动性"), ("风控", "波动/尾部风险")]
    for i, (role, theme) in enumerate(roles):
        x = 55 + i * 265
        helpers.rounded(d, (x, 130, x + 235, 250), "#ffffff", outline=helpers.COLORS[i], radius=14, width=3)
        d.text((x + 117, 168), f"{role}分析师", font=helpers.font(22, True), fill="#203c50", anchor="ma")
        d.text((x + 117, 215), f"8 因子 · {theme}", font=helpers.font(17), fill="#607787", anchor="ma")
    d.text((700, 292), "默认初始化：每个角色内部 8 个因子各 12.5%（透明基线，不代表最优）", font=helpers.font(21, True), fill="#36566d", anchor="ma")
    boxes = [
        (70, 360, 325, 575, "客户意图", "快捷标签或自由文字\n例如：更稳健、降低波动"),
        (395, 360, 650, 575, "结构化建议", "仅允许白名单角色、因子\n及模块参数进入补丁"),
        (720, 360, 975, 575, "安全约束", "角色内归一化 100%\n单因子上限 35%\n风险因子按方向处理"),
        (1045, 360, 1300, 575, "人工确认", "展示调整前后、理由、来源\n确认后生成 config_id"),
    ]
    for i, (x1, y1, x2, y2, title, body) in enumerate(boxes):
        helpers.rounded(d, (x1, y1, x2, y2), "#ffffff", outline="#a8bac6", radius=16, width=3)
        d.text(((x1 + x2) // 2, y1 + 45), title, font=helpers.font(25, True), fill=helpers.COLORS[i], anchor="ma")
        for j, line in enumerate(body.split("\n")):
            d.text((x1 + 20, y1 + 105 + j * 42), line, font=helpers.font(17), fill="#526b7c")
        if i < len(boxes) - 1:
            helpers.arrow(d, x2 + 10, (y1 + y2) // 2, boxes[i + 1][0] - 10)
    helpers.rounded(d, (70, 640, 1300, 730), "#fff5e6", outline="#d8bd8b", radius=14, width=2)
    d.text((685, 675), "DeepSeek 失败、超时、JSON 错误或越界时明确标记 rule_fallback；信息不足时先澄清，不修改配置。", font=helpers.font(20, True), fill="#6f5422", anchor="ma")
    img.save(out, quality=95)
    return out


def save_evidence_flow() -> Path:
    out = ASSET_DIR / "fig5-9_evidence_explain_v04.png"
    img = Image.new("RGB", (1400, 750), "#f6f8fa")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "回测证据白名单与前端可解释链路", font=helpers.font(40, True), fill="#19324a")
    steps = [
        ("原始回测结果", "收益、风险、角色贡献\nrequest_id / config_id"),
        ("一致性校验", "有限数、收益恒等式\n权重合计、50% 上限\nShapley 可加性"),
        ("可信证据白名单", "合法字段生成 E001…\n异常字段隔离并告警"),
        ("Agent 解读", "引用合法证据编号\n无引用或越界则降级"),
        ("数据依据抽屉", "定义、公式、本次代入\n区间、基准、字段与局限"),
    ]
    x0, y0, w, h, gap = 35, 165, 240, 330, 35
    for i, (title, body) in enumerate(steps):
        x = x0 + i * (w + gap)
        helpers.rounded(d, (x, y0, x + w, y0 + h), "#ffffff", outline="#a8bac6", radius=15, width=3)
        d.text((x + 20, y0 + 47), title, font=helpers.font(22, True), fill=helpers.COLORS[i])
        cursor = 130
        for line in body.split("\n"):
            for wrapped in fit_text(d, line, w - 40, 17):
                d.text((x + 20, y0 + cursor), wrapped, font=helpers.font(17), fill="#526b7c")
                cursor += 39
        if i < len(steps) - 1:
            helpers.arrow(d, x + w + 7, y0 + h // 2, x + w + gap - 7)
    helpers.rounded(d, (65, 570, 1335, 690), "#e9f5f0", outline="#a8ccbd", radius=15, width=2)
    d.text((700, 610), "前端每个核心指标可追溯到后端字段和证据编号", font=helpers.font(22, True), fill="#285b4b", anchor="ma")
    d.text((700, 650), "累计收益 R = ∏(1+r_t)-1；超额收益 = 组合累计收益 - 同期 HS300 累计收益", font=helpers.font(19), fill="#36566d", anchor="ma")
    img.save(out, quality=95)
    return out


def add_picture_before(target, path: Path, width: float, alt: str):
    paragraph = target.insert_paragraph_before()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt)
    return paragraph


def add_table_before(doc: Document, target, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    helpers.set_table(table, headers, rows)
    target._p.addprevious(table._tbl)
    return table


def main():
    overview_chart = save_overview_flow()
    factor_chart = save_factor_agent_flow()
    evidence_chart = save_evidence_flow()

    doc = Document(SOURCE)
    tech_table = doc.tables[0]
    dual_track_table = doc.tables[1]
    api_table = doc.tables[2]
    milestone_table = doc.tables[4]
    appendix_figure_table = doc.tables[5]
    replace_by_prefix(doc, "版本：", "版本：v0.4（对齐客户专属因子 Agent、证据约束、四步自助流程与 85 项回归测试）")
    replace_by_prefix(doc, "本分册覆盖技术文档", "本分册覆盖技术文档分工中高玉伟主笔/收口内容。系统以行业、个股、策略、量化、风控五类投研角色为直接计量对象，每个角色关联 8 个可调整因子，共形成 40 因子客户专属方案。用户可通过快捷选择和自然语言向 Agent 描述需求；Agent 只生成受白名单和边界约束的配置建议，人工确认后方可进入回测。主线仍以 32 个角色联盟的反事实回测计算 Shapley 超额贡献，并通过制度轨 60% + 证据轨 40% 形成广义价值分配建议。")
    replace_by_prefix(doc, "本文中的架构图", "本文中的架构图与流程图用于解释从客户意图、因子配置、回测证据到人工确认的技术闭环；第5.6节数值表和 Stage-3 结果图取自当前完整真实回测，作为定量结论依据。客户专属配置能力用于解释方案差异，不构成收益承诺。")
    replace_by_prefix(doc, "AlphaSplit 构建", "AlphaSplit 构建“多智能体驱动的投研角色贡献计量平台”：以五类投研角色为计量对象，将每个角色的 8 个决策因子组织为可审计配置。客户通过快捷选择或自然语言描述投资偏好，DeepSeek 因子 Agent 将意图转换为白名单结构化补丁，经约束校验和人工确认后进入客户专属回测。系统在相同日期、基准和调仓频率下对照客户方案、默认等权方案与沪深300，再通过 32 个角色联盟的反事实回测计算 Shapley 超额贡献，最后进入 6:4 双轨和人工确认流程。")
    replace_by_prefix(doc, "当前工程已形成", "当前工程已形成可本地运行闭环：FastAPI + ECharts Web 页面、40 因子客户配置、DeepSeek/规则降级 Agent、HS300 实时或缓存回测、SSE 进度流、回测证据白名单、数据依据抽屉、双轨合成与五角色协商。Stage-3 回测覆盖 2024-01-01 至 2024-06-30、297 只股票、117 个交易日和 32 个联盟：组合累计收益 29.51%，同期 HS300 收益 7.66%，超额收益及 Shapley 合计均为 21.85%，有效性校验通过，联盟失败数为 0。")
    replace_by_prefix(doc, "系统按信号", "系统按信号、计量、智能和治理输出四层组织。信号层沉淀行情、行业映射、五角色 40 因子与客户配置；计量层执行候选池构建、客户/默认/基准对照、组合回测及全联盟 Shapley 超额贡献归因；智能层负责意图解释、配置建议、证据约束解读与多角色协商；治理输出层形成制度轨/证据轨双轨方案，并要求人工确认。")
    replace_by_prefix(doc, "（1）信号层", "（1）信号层：data/processed/ 下 HS300 日线、行业映射和财务数据，以及 backtest/factor_config.py 中的五角色 40 因子登记、方向、默认权重和参数边界。")
    replace_by_prefix(doc, "（3）智能层", "（3）智能层：agent_layer/ 负责 NL→回测计划、客户因子建议与证据约束解读；engine/agentic/role_negotiation.py 实现五角色协商。DeepSeek 的 API Key 仅从服务端环境变量 DEEPSEEK_API_KEY 读取。")
    replace_by_prefix(doc, "仓库关键模块包括", "仓库关键模块包括 engine/api（FastAPI 与访问保护）、engine/web（四步自助流程、ECharts 与数据依据抽屉）、engine/integrations（回测与 Agent 桥接）、backtest（40 因子配置、五角色 Pipeline 与 Shapley 引擎）、agent_layer（规划、因子建议、证据白名单与解读）、engine/agentic（协商与附录 Mock）以及 data/scenarios。依赖由 requirements-lock.txt 锁定，安装与前台启动职责分离。")
    replace_by_prefix(doc, "用户/Agent 规划", "客户描述需求 → Agent 生成白名单因子补丁 → 约束校验与人工确认 → 传递 factor_config → 客户方案/默认方案/HS300 同口径对照 → 五角色 Pipeline 回测 → 31 个非空联盟 + 空联盟反事实 → Shapley 超额贡献 → 证据白名单 → 6:4 双轨合成 → 五角色协商 → 最终人工确认。")
    replace_by_prefix(doc, "可运行性与安全性", "可运行性与安全性：官方演示包可离线秒开；因子建议支持 DeepSeek 与明确标记的 rule_fallback；LLM 采用有限重试，JSON Schema 不支持时降级为 JSON Object，不从自然语言中正则猜测 JSON。公开演示模式可关闭历史结果与真实回测接口；配置 API Token 后，真实回测、Agent Pipeline 和历史结果均要求 X-AlphaSplit-Token。")
    replace_by_prefix(doc, "2.3–2.5 由", "2.3–2.5 由董斌/徐远主笔。统稿时需保证五角色、40 因子与第5章输出维度一致，并明确 12.5% 是角色内等权初始化基线而非经验最优值。JSON 中 role_weights、old_pool、new_pool 等历史字段为兼容接口，正文统一称制度调节系数、制度轨和证据轨。")
    replace_by_prefix(doc, "人工确认页由", "人工确认页由 engine/agentic/role_negotiation.py 驱动五角色协商。协商输入是已经包含 60% 原有方案和 40% 纯新方案的最终合成锚点；发言中的份额指总可分配资源池份额，而非仅指 40% 证据轨。各角色引用 Shapley 超额贡献、纯新方案比例和最终合成锚点陈述主张，治理责任人决定是否采纳。")
    replace_by_prefix(doc, "回测规划与解读", "回测规划、因子建议与解读由独立 agent_layer 实现，并经 engine/integrations/agent_bridge.py 作为唯一桥接入口。因子 Agent 将客户意图转换为结构化配置补丁；信息不足时澄清而不修改，DeepSeek 失败、超时、JSON 错误或越界时明确标记 rule_fallback。建议必须由用户确认后才进入回测。")
    replace_by_prefix(doc, "截至本版", "截至本版，85 项单元与接口测试全部通过，前端 JavaScript 语法检查和 Git 差异格式检查通过；桌面端与 390×844 移动端已完成浏览器流程验证，页面错误日志为空。测试覆盖 40 因子归一化与上限、LLM 重试与安全降级、回测证据一致性、四步流程锁定、下一步自动定位、数据依据抽屉和 6:4 比例口径。")
    replace_by_prefix(doc, "（4）Agent 默认", "（4）客户因子 Agent 支持 DeepSeek 和明确标记的规则降级；任何配置或应用建议均须人工确认。")
    replace_by_prefix(doc, "当前局限包括", "当前局限包括：角色因子仍是岗位决策代理；每角色 12.5% 等权仅是透明初始化基线；单一 2024H1 区间不能代表所有市场状态；客户专属方案仍需更多滚动窗口和样本外验证；真实回测为同步长任务；机构脱敏决策日志仍待获取。")
    replace_by_prefix(doc, "后续将开展", "后续将开展滚动窗口与样本外验证，比较客户方案、默认方案与基准的稳定性，报告贡献均值、正贡献比例和置信区间；同时评估任务队列、取消/超时机制、配置版本迁移、Agent 质量评测及端到端浏览器持续回归。")

    first_image_paragraph = doc.paragraphs[20]
    helpers.replace_image(first_image_paragraph, overview_chart, 6.15)
    first_shape = doc.inline_shapes[0]
    first_shape._inline.docPr.set("descr", "客户描述需求、确认Agent因子方案、完成三方回测并形成应用建议的四步闭环")
    first_shape._inline.docPr.set("title", "AlphaSplit客户专属方案四步技术闭环")

    helpers.set_table(tech_table, ["类别", "组件", "用途"], [
        ["语言", "Python 3.11（已验证）", "回测、API 与 Agent 主业务"],
        ["API", "FastAPI / Uvicorn", "结构化 /api/*、SSE 与访问保护"],
        ["回测", "backtest/query_engine.py", "客户配置、五角色 Pipeline 与超额收益 Shapley"],
        ["因子配置", "backtest/factor_config.py", "40 因子登记、方向、默认值、归一化与边界"],
        ["前端", "HTML/CSS/JS + ECharts", "四步自助流程、三方对照、数据依据抽屉与人工确认"],
        ["Agent", "agent_layer（DeepSeek / rule_fallback）", "规划、因子建议、证据白名单与结果解读"],
        ["数据", "data/processed/*.csv", "HS300 行情、行业映射与因子"],
        ["工程", "requirements-lock.txt + tests", "依赖锁定、接口保护与 85 项回归验证"],
    ])
    helpers.set_table(dual_track_table, ["比例口径", "占比", "含义", "输入/来源"], [
        ["原有方案比例", "制度轨内部 100%", "仅描述 60% 制度轨内部结构", "用户录入；演示默认五角色各 20%"],
        ["纯新方案比例", "证据轨内部 100%", "仅描述 40% 证据轨内部结构", "role_breakdown.weight"],
        ["最终合成比例", "总资源池 100%", "60% 原有方案 + 40% 纯新方案", "协商锚点与人工确认输入"],
    ])
    helpers.set_table(api_table, ["接口", "作用"], [
        ["GET /api/agent/factor-config", "返回五角色、40 因子元数据和默认配置"],
        ["POST /api/agent/factor-config", "根据客户意图返回受约束的配置变更预览"],
        ["POST /api/agent/plan", "自然语言或预设生成结构化回测计划"],
        ["POST /api/agent/interpret", "基于可信证据编号解读回测结果"],
        ["POST /api/agent/pipeline", "串联 Agent 规划、回测与证据约束解读"],
        ["POST /api/backtest/run", "使用已确认 factor_config 执行五角色回测与 Shapley"],
        ["POST /api/backtest/run/stream", "SSE 推送阶段、联盟与完成事件"],
        ["GET /api/backtest/demo", "读取 Stage-3 官方预置结果"],
        ["POST /api/negotiate", "五角色基于证据和最终合成锚点协商"],
        ["POST /api/confirm", "人工确认最终应用建议"],
    ])
    helpers.set_table(milestone_table, ["阶段", "事项", "产出"], [
        ["工程骨架", "FastAPI + Web Demo", "可本地前台启动"],
        ["归因修正", "超额收益 Shapley + 失败中止", "32 联盟有效性校验"],
        ["角色链路", "个股候选池 + 量化独立优化", "角色职责可检验"],
        ["客户方案", "五角色 × 8 因子配置", "可版本化 factor_config"],
        ["因子 Agent", "DeepSeek + rule_fallback + 人工确认", "受约束配置建议"],
        ["三方对照", "客户 / 默认 / HS300", "同口径解释配置差异"],
        ["可信解释", "证据白名单 + E001 编号", "异常隔离与可追溯解读"],
        ["前端流程", "四步解锁 + 下一步自动定位", "首次用户可自助完成"],
        ["双轨治理", "三种比例口径 + 五角色协商", "广义价值分配建议"],
        ["测试回归", "85 项 + JS/格式/浏览器验证", "桌面与移动端通过"],
        ["文档 v0.4", "代码逻辑、口径和验证同步", "本完善版"],
    ])

    chapter7 = paragraph_by_prefix(doc, "第7章供料")

    helpers.add_before(chapter7, "5.9 客户专属因子方案与 DeepSeek Agent", style="Heading 2")
    helpers.add_before(chapter7, "客户专属投资方案由五角色的 40 个因子及少量模块参数共同描述。行业、个股、策略、量化和风控分析师各有 8 个因子；默认配置 balanced_equal_v2 在每个角色内部采用 12.5% 等权，仅作为透明、可复现的初始化基线，不表示已经通过样本外验证的最优权重。配置进入回测前必须角色内归一化为 100%，单因子最终权重不得超过 35%。波动率、ATR、下行波动等风险因子采用低值优先方向，避免“风险越高得分越高”的符号错误。")
    add_table_before(doc, chapter7, ["角色", "8 个因子（中文标签）", "默认与方向口径"], [
        ["行业分析师", "5日动量、10日动量、均线偏离、均线斜率、趋势强度、趋势效率、突破信号、区间位置", "角色内各12.5%；高值优先"],
        ["个股分析师", "5日反转、10日反转、RSI强度、RSI反转、随机指标、回撤、下影线支撑、下轨突破", "角色内各12.5%；基础反向已在引擎处理"],
        ["策略分析师", "跳过期间动量、跳空溢价、日内收益、已实现波动率、ATR比率、区间波动比、收益偏度、收益峰度", "波动、ATR、区间波动和峰度低值优先"],
        ["量化分析师", "成交量比率、成交量Z分数、成交额活跃度、量价相关性、OBV斜率、收盘时序排名、成交量时序排名、下行波动率", "下行波动率低值优先"],
        ["风控分析师", "5日/10日已实现波动率、下行波动率、ATR比率、区间波动比、回撤、收益偏度、收益峰度", "多数风险指标低值优先；回撤越接近0越好"],
    ])
    add_picture_before(chapter7, factor_chart, 6.0, "客户意图经DeepSeek或规则降级转换为受约束的40因子配置并由人工确认")
    helpers.add_before(chapter7, "图5-8  客户专属 40 因子方案与受约束 Agent 调整闭环")
    helpers.add_before(chapter7, "交互采用“快捷选择 + 自由文字”的混合方式。Agent 输出只允许包含登记过的角色、因子和模块参数；信息不足时返回澄清问题，不修改当前配置。DeepSeek 调用失败、超时、JSON 错误或越界时，系统明确返回 rule_fallback 和原因，不把规则结果伪装成模型结果。API Key 仅保存在服务端 DEEPSEEK_API_KEY 环境变量中。")
    helpers.add_before(chapter7, "可靠性方面，LLM 采用有限重试；若服务端不支持 JSON Schema，则降级到 JSON Object 响应格式，但不会从自由文本中用正则猜测 JSON。每次建议记录独立审计号、配置前后指纹、变更数量、来源和人工确认状态；日志不记录 API Key、完整提示词、完整服务地址或原始错误正文。")

    helpers.add_before(chapter7, "5.10 客户专属回测、三方对照与证据约束", style="Heading 2")
    helpers.add_before(chapter7, "人工确认后的 factor_config 随回测请求传入后端，只替换被修改的因子或模块参数，其他配置保持不变。页面在相同日期范围、HS300 基准和调仓频率下对照客户方案、默认等权方案和沪深300，比较累计收益、超额收益、年化夏普比率与最大回撤。该对照用于解释配置差异，不用于反复样本内调参，也不构成未来收益承诺。")
    helpers.add_before(chapter7, "回测结果进入解读 Agent 前先建立只读证据白名单：检查所有数值是否有限、最大回撤范围、超额收益恒等式、角色权重合计、单角色 50% 上限和 Shapley 可加性。不一致字段被隔离并生成结构化警告；可信字段生成 E001 等证据编号。LLM 涉及数值的结论必须引用合法证据编号，出现未知编号、无证据数字或收益承诺时自动降级为确定性解读。")
    add_picture_before(chapter7, evidence_chart, 6.0, "原始回测结果经一致性校验形成E001证据白名单并支持前端可解释追溯")
    helpers.add_before(chapter7, "图5-9  回测证据白名单、Agent 解读与前端数据依据链路")
    helpers.add_before(chapter7, "前端“数据依据”抽屉为每个核心指标展示定义、公式、本次代入、时间范围、基准、后端字段、证据编号、request_id、config_id 和使用局限。桌面端使用右侧抽屉，移动端使用底部面板。累计收益按 R=∏(1+r_t)-1 计算；超额收益为组合累计收益减同期 HS300 累计收益；Shapley 贡献对应各角色对完整超额收益的平均边际贡献。")

    helpers.add_before(chapter7, "5.11 前端四步自助流程与交互治理", style="Heading 2")
    helpers.add_before(chapter7, "首次用户按四个步骤完成闭环：（1）描述投资需求；（2）确认 Agent 方案；（3）回测并理解结果；（4）形成应用建议。未来步骤在前置动作完成前锁定；当前唯一下一步按钮通过状态标签和视觉提示高亮，并在生成建议、确认配置、完成回测和进入应用建议后自动定位。页面刷新时关闭浏览器旧滚动位置恢复，并定位到当前有效的下一步按钮。")
    helpers.add_before(chapter7, "第3步默认推荐“使用预置数据快速体验”，使首次用户无需等待完整实时回测即可理解系统；“运行我的方案回测”保留为验证已确认客户因子配置的可选操作。高级参数、缓存和技术日志默认收纳在专业模式，减少首次使用者的认知负担。")
    helpers.add_before(chapter7, "第4步明确区分三种比例：原有方案比例只描述制度轨内部结构；纯新方案比例只描述 40% 证据轨内部结构；最终合成比例占总可分配资源池。五角色协商和人工确认处理的是已经包含 60% 原有方案与 40% 纯新方案的最终比例，页面不得把它误写成“40% 新方案内部比例”。")
    helpers.add_before(chapter7, "5.12 当前验证状态与适用边界", style="Heading 2")
    helpers.add_before(chapter7, "完整测试共 85 项全部通过；node --check engine/web/app.js 通过；git diff --check 无补丁格式错误。浏览器已验证桌面端和 390×844 移动端的四步流程、步骤锁定、下一步自动定位、预置回测解锁、数据依据面板、三种比例口径和最终人工确认，页面错误日志为空。")
    helpers.add_before(chapter7, "适用边界：系统输出的是客户方案回测、角色贡献证据和广义价值分配建议，不替代投资决策、风险审批或治理责任人的判断。默认因子权重和单一区间结果均需通过滚动窗口、样本外数据和稳定性分析继续验证。")

    helpers.set_table(appendix_figure_table, ["图号", "文件", "用途"], [
        ["图1-1", overview_chart.name, "客户专属方案四步技术闭环"],
        ["图2-1", "fig2-1_architecture_stage3.png", "总体架构"],
        ["图2-2", "fig2-2_modules_stage3.png", "模块关系"],
        ["图2-3", "fig2-3_web_flow_stage3.png", "Web 工作流"],
        ["图4-3", "fig4-3_mock.png", "Mock/LLM 模式"],
        ["图5-1", "fig5-1_shapley_stage3.png", "Shapley 机制"],
        ["图5-3", "fig5-3_loo_shapley_stage3.png", "LOO 方法对照"],
        ["图5-4", "fig5-4_dual_track_stage3.png", "双轨合成"],
        ["图5-5", "fig5-4e_agent_panel.png", "Agent 规划"],
        ["图5-6", "fig5-6_negotiation_flow.png", "协商与人工确认"],
        ["图5-7", "fig5-7_stage3_results.png", "Stage-3 回测结果"],
        ["图5-8", factor_chart.name, "40 因子 Agent 调整闭环"],
        ["图5-9", evidence_chart.name, "证据白名单与可解释链路"],
    ])

    full_text = "\n".join(p.text for p in doc.paragraphs)
    full_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    found = [term for term in NARROW_TERMS if term in full_text]
    if found:
        raise RuntimeError(f"narrow allocation terms remain: {found}")
    if "53 项" in full_text or "点击级浏览器自动化尚未接入" in full_text:
        raise RuntimeError("outdated verification statement remains")

    doc.save(OUTPUT)
    print(f"saved {OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
